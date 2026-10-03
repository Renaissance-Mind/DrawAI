#!/usr/bin/env python3
"""Exercise the documented document-only workflows with real files."""

from pathlib import Path
import json
import shutil
import hashlib
import subprocess
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lxml import etree
from pptx import Presentation
from pptx_skill import (
    Document,
    export_pptx,
    export_svg,
    import_pptx,
    import_svg,
    duplicate_page,
    inspect_pptx,
    make_connector,
)
from pptx_skill.document.identity import IDENTITY_NS

HERE = Path(__file__).resolve().parent


def box(id, kind="shape", **properties):
    return {
        "id": id,
        "kind": kind,
        "properties": {
            "object.name": id,
            "object.frame": {"left": 20, "top": 20, "width": 120, "height": 80},
            **properties,
        },
    }


def document(objects):
    return Document.model_validate(
        {
            "schema": "presentation.document.v1",
            "id": "deck",
            "properties": {"document.page_size": {"width": 720, "height": 405}},
            "pages": [{"id": "page", "objects": objects}],
        }
    )


class Workflows(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_crop_full_keys_short_keys_and_import_contract(self):
        doc = document(
            [
                box(
                    "photo",
                    "picture",
                    **{
                        "picture.asset_ref": "image",
                        "picture.crop": {"left": 0.1, "t": 0.2, "right": 0.05},
                    },
                )
            ]
        )
        doc.assets = {
            "image": {"path": str(HERE.parent / "assets/demo/demo_figure.png")}
        }
        p = self.root / "crop.pptx"
        export_pptx(doc, p)
        s = Presentation(p).slides[0].shapes[0]
        self.assertAlmostEqual(s.crop_left, 0.1)
        self.assertAlmostEqual(s.crop_top, 0.2)
        self.assertEqual(
            set(import_pptx(p).pages[0].objects[0].properties["picture.crop"]),
            {"left", "top", "right", "bottom"},
        )

    def test_identity_aliases_create_and_conflicts_reject(self):
        raw = {"kind": "shape", "id": "a", "properties": {"object.id": "b"}}
        with self.assertRaisesRegex(ValueError, "conflicts"):
            document([raw])
        d = document(
            [
                {
                    "properties": {
                        "object.id": "text",
                        "object.kind": "text_box",
                        "object.frame": {
                            "left": 10,
                            "top": 10,
                            "width": 200,
                            "height": 50,
                        },
                        "text.body.paragraphs": [{"content": [{"text": "Title"}]}],
                    }
                }
            ]
        )
        p = self.root / "alias.pptx"
        export_pptx(d, p)
        self.assertEqual(Presentation(p).slides[0].shapes[0].shape_type, 17)
        d.pages[0].objects[0].properties["object.id"] = "different"
        with self.assertRaisesRegex(ValueError, "conflicts"):
            export_pptx(d, self.root / "bad.pptx")
        self.assertFalse((self.root / "bad.pptx").exists())

    def test_clear_format_survives_text_reconstruction(self):
        d = document(
            [
                box(
                    "title",
                    "text_box",
                    **{
                        "text.body.paragraphs": [
                            {
                                "content": [
                                    {
                                        "text": "Original",
                                        "properties": {"text.run.font.bold": True},
                                    }
                                ]
                            }
                        ]
                    },
                )
            ]
        )
        p = self.root / "bold.pptx"
        export_pptx(d, p)
        d = import_pptx(p)
        o = d.pages[0].objects[0]
        o.clear_overrides = ["text.run.font.bold"]
        o.properties["text.body.paragraphs"][0]["content"][0]["text"] = "Updated"
        q = self.root / "updated.pptx"
        export_pptx(d, q)
        r = Presentation(q).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        self.assertEqual(r.text, "Updated")
        self.assertIsNone(r.font.bold)

    def test_existing_page_and_group_stacking_order(self):
        group = box("group", "group")
        group["children"] = [box("child-back"), box("child-front")]
        p = self.root / "layers.pptx"
        export_pptx(document([box("back"), box("front"), group]), p)
        d = import_pptx(p)
        d.pages[0].objects[0].properties["object.z_order"] = 3
        d.pages[0].objects[1].properties["object.z_order"] = 1
        d.pages[0].objects[2].properties["object.z_order"] = 2
        d.pages[0].objects[2].children[0].properties["object.z_order"] = 2
        d.pages[0].objects[2].children[1].properties["object.z_order"] = 1
        q = self.root / "layers-updated.pptx"
        export_pptx(d, q)
        s = Presentation(q).slides[0].shapes
        self.assertEqual([o.name for o in s], ["front", "group", "back"])
        self.assertEqual([o.name for o in s[1].shapes], ["child-front", "child-back"])

    def test_copy_and_import_duplicate_extension_ids(self):
        p = self.root / "source.pptx"
        export_pptx(document([box("shape")]), p)
        d = duplicate_page(import_pptx(p), "page", new_id="copy")
        q = self.root / "copied.pptx"
        export_pptx(d, q)
        self.assertEqual(
            len({o.id for page in import_pptx(q).pages for o in page.objects}), 2
        )
        with zipfile.ZipFile(q) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        first = etree.fromstring(parts["ppt/slides/slide1.xml"])
        second = etree.fromstring(parts["ppt/slides/slide2.xml"])
        for a, b in zip(
            first.iter("{" + IDENTITY_NS + "}identity"),
            second.iter("{" + IDENTITY_NS + "}identity"),
            strict=True,
        ):
            b.set("id", a.get("id"))
        parts["ppt/slides/slide2.xml"] = etree.tostring(second)
        modified = self.root / "editor-copy.pptx"
        with zipfile.ZipFile(modified, "w") as z:
            for n, b in parts.items():
                z.writestr(n, b)
        restored = import_pptx(modified)
        self.assertEqual(len({p.id for p in restored.pages}), 2)
        self.assertEqual(len({o.id for p in restored.pages for o in p.objects}), 2)

    @unittest.skipUnless(shutil.which("soffice"), "SVG projection needs LibreOffice")
    def test_page_svg_recovers_full_source_with_page_bindings(self):
        d = document([box("a")])
        raw = d.model_dump(by_alias=True)
        raw["pages"].append({"id": "page-B", "objects": [box("b")]})
        raw["native_bindings"] = [
            {
                "owner_id": "page-B",
                "property_id": "slide.name",
                "part": "ppt/slides/slide2.xml",
                "xpath": "/p:sld/p:cSld",
                "action": "set_attribute",
                "name": "name",
                "value": "Bound B",
            }
        ]
        d = Document.model_validate(raw)
        r = export_svg(d, self.root / "svg")
        recovered, report = import_svg(Path(r["pages"][0]["svg"]))
        self.assertEqual(
            recovered.model_dump(by_alias=True), d.model_dump(by_alias=True)
        )
        self.assertEqual(report["source_scope"], "full_document")
        self.assertEqual(report["projection_page_id"], "page")
        p = self.root / "recovered.pptx"
        export_pptx(recovered, p)
        self.assertEqual(Presentation(p).slides[1]._element.cSld.get("name"), "Bound B")

    def test_cli_retains_svg_report_and_queries_properties(self):
        svg = HERE.parent / "assets/demo/formula.svg"
        out = self.root / "import.json"
        report = self.root / "import-report.json"
        run = subprocess.run(
            [
                sys.executable,
                "-I",
                str(HERE / "ppt.py"),
                "import",
                str(svg),
                "--output",
                str(out),
                "--report",
                str(report),
            ],
            capture_output=True,
            text=True,
            check=True,
        )
        result = json.loads(run.stdout)
        self.assertIn("backend", result["conversion"])
        self.assertEqual(json.loads(report.read_text()), result)
        run = subprocess.run(
            [
                sys.executable,
                "-I",
                str(HERE / "ppt.py"),
                "capabilities",
                "--property",
                "picture.crop",
            ],
            capture_output=True,
            text=True,
            check=True,
        )
        self.assertEqual(
            [p["id"] for p in json.loads(run.stdout)["properties"]], ["picture.crop"]
        )

    def test_complete_example_create_and_edit(self):
        create = self.root / "create"
        edit = self.root / "edit"
        subprocess.run(
            [sys.executable, "-I", str(HERE / "example.py"), str(create)],
            capture_output=True,
            text=True,
            check=True,
        )
        subprocess.run(
            [
                sys.executable,
                "-I",
                str(HERE / "example.py"),
                str(edit),
                "--source",
                str(create / "deck.pptx"),
                "--object-id",
                "heading",
                "--find",
                "Original ",
                "--text",
                "Updated ",
            ],
            capture_output=True,
            text=True,
            check=True,
        )
        runs = (
            Presentation(edit / "deck.pptx")
            .slides[0]
            .shapes[0]
            .text_frame.paragraphs[0]
            .runs
        )
        self.assertEqual(runs[0].text, "Updated ")
        self.assertTrue(runs[0].font.bold)
        self.assertTrue(runs[1].font.italic)

    def test_layout_placeholder_without_geometry_inspects(self):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        for owner in (prs.slide_masters[0], prs.slide_layouts[0]):
            for placeholder in owner.placeholders:
                if placeholder._element.xpath("./p:nvSpPr/p:nvPr/p:ph[@type='dt']"):
                    for xfrm in placeholder._element.xpath("./p:spPr/a:xfrm"):
                        xfrm.getparent().remove(xfrm)
        path = self.root / "inherited.pptx"
        prs.save(path)
        result = inspect_pptx(path)
        dates = [
            p
            for p in result["layouts"][0]["placeholders"]
            if p["type"].startswith("DATE")
        ]
        self.assertEqual(len(dates), 1)
        self.assertTrue(all(v is None for v in dates[0]["frame"].values()))

    def test_connectors_preserve_all_endpoint_directions(self):
        start = (140, 140)
        ends = [
            (30, 30),
            (250, 30),
            (30, 250),
            (250, 250),
            (30, 140),
            (250, 140),
            (140, 30),
            (140, 250),
        ]
        raw = document([]).model_dump(by_alias=True)
        raw["pages"][0]["objects"] = [
            make_connector(
                f"line-{i}",
                start,
                end,
                properties={"stroke.end_arrow.kind": "triangle"},
            ).model_dump()
            for i, end in enumerate(ends)
        ]
        path = self.root / "directions.pptx"
        export_pptx(raw, path)
        for shape, end in zip(Presentation(path).slides[0].shapes, ends, strict=True):
            self.assertEqual(
                (shape.begin_x, shape.begin_y), tuple(round(v * 12700) for v in start)
            )
            self.assertEqual(
                (shape.end_x, shape.end_y), tuple(round(v * 12700) for v in end)
            )
            self.assertEqual(
                shape._element.xpath("./p:spPr/a:ln/a:tailEnd")[0].get("type"),
                "triangle",
            )
        bad = box("reverse", "connector")
        bad["properties"]["object.frame"]["width"] = -100
        with self.assertRaisesRegex(ValueError, "Object reverse"):
            document([bad])

    @unittest.skipUnless(
        shutil.which("fc-match"),
        "An installed font is needed for the embedded-font package",
    )
    def test_render_font_copy_preserves_original_and_slide_parts(self):
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from pptx.opc.package import Part
        from pptx_skill.document.render_env import office_input

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        font = Path(
            subprocess.check_output(
                ["fc-match", "-f", "%{file}", "sans-serif"], text=True
            )
        )
        part = Part(
            prs.part.package.next_partname("/ppt/fonts/font%d.fntdata"),
            "application/x-fontdata",
            prs.part.package,
            font.read_bytes(),
        )
        rid = prs.part.relate_to(part, RT.FONT)
        fonts = etree.SubElement(
            prs._element,
            "{http://schemas.openxmlformats.org/presentationml/2006/main}embeddedFontLst",
        )
        entry = etree.SubElement(
            fonts,
            "{http://schemas.openxmlformats.org/presentationml/2006/main}embeddedFont",
        )
        etree.SubElement(
            entry,
            "{http://schemas.openxmlformats.org/presentationml/2006/main}font",
            typeface="sans-serif",
        )
        etree.SubElement(
            entry,
            "{http://schemas.openxmlformats.org/presentationml/2006/main}regular",
            {
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": rid
            },
        )
        source = self.root / "embedded.pptx"
        prs.save(source)
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        render_source, report = office_input(source, self.root / "temporary")
        self.assertTrue(report["rendering_copy_used"])
        self.assertEqual(report["embedded_font_families"], ["sans-serif"])
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
        with zipfile.ZipFile(source) as a, zipfile.ZipFile(render_source) as b:
            for name in a.namelist():
                if name not in (
                    "ppt/presentation.xml",
                    "ppt/_rels/presentation.xml.rels",
                ):
                    self.assertEqual(a.read(name), b.read(name), name)
            self.assertNotIn(b"embeddedFontLst", b.read("ppt/presentation.xml"))
        from pptx_skill import validate_pptx_package

        self.assertEqual(validate_pptx_package(render_source)["status"], "ok")
        kept, result = office_input(source, self.root / "preserve", "preserve")
        self.assertEqual(kept, source)
        self.assertFalse(result["rendering_copy_used"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
