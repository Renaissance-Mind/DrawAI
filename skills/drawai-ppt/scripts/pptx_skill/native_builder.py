"""Assemble native SVG drawings into a self-contained PPTX package."""

from __future__ import annotations

from io import BytesIO
import json
import math
from pathlib import Path
import zipfile

from lxml import etree
from pptx import Presentation

from .svg_native.drawingml_converter import convert_svg_to_slide_shapes


EMU_PER_PX = 9525
_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_MEDIA_TYPES = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                "gif": "image/gif", "bmp": "image/bmp", "tif": "image/tiff",
                "tiff": "image/tiff", "svg": "image/svg+xml", "emf": "image/x-emf",
                "wmf": "image/x-wmf", "webp": "image/webp"}


def create_pptx_with_native_svg(svg_files, output_path, *, canvas_format=None,
                              verbose=False, transition=None, use_native_shapes=True,
                              enable_notes=False, animation=None, merge_paragraphs=True,
                              conversion_trace_path=None, doc_metadata=None):
    """Export one SVG fragment with actual layout and media relationships."""
    if len(svg_files) != 1 or not use_native_shapes or enable_notes or animation or transition or canvas_format:
        raise ValueError("The fragment exporter accepts one native SVG without slide playback settings")
    source = Path(svg_files[0])
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    root = etree.fromstring(source.read_bytes(), parser)
    viewbox = root.get("viewBox", "").replace(",", " ").split()
    if len(viewbox) != 4:
        raise ValueError("SVG fragment must declare a four-number viewBox")
    _, _, width, height = map(float, viewbox)
    if not all(math.isfinite(value) and value > 0 for value in (width, height)):
        raise ValueError("SVG viewBox dimensions must be finite and positive")
    prs = Presentation()
    scale = max(1.0, 96 / width, 96 / height)
    if max(width, height) * scale > 5376:
        scale = 5376 / max(width, height)
    if min(width, height) * scale < 96:
        raise ValueError("SVG aspect ratio cannot fit PowerPoint's 1–56 inch slide size limits")
    prs.slide_width, prs.slide_height = round(width * scale * EMU_PER_PX), round(height * scale * EMU_PER_PX)
    # This is the standard blank layout in python-pptx's bundled default deck.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    metadata = doc_metadata or {}
    prs.core_properties.title = metadata.get("title", Path(output_path).stem)
    prs.core_properties.subject = metadata.get("subject", "Native SVG drawing")
    prs.core_properties.author = "drawai-ppt"
    buffer = BytesIO()
    prs.save(buffer)
    with zipfile.ZipFile(buffer) as base:
        contents = {name: base.read(name) for name in base.namelist()}
    traces = []
    drawing, media, relationships, _ = convert_svg_to_slide_shapes(
        source, slide_num=1, verbose=verbose, merge_paragraphs=merge_paragraphs, trace_out=traces)
    if scale != 1:
        from .document.xml import el, tag
        page = etree.fromstring(drawing.encode(), parser)
        tree = page.find('.//' + tag('p:spTree'))
        content = [n for n in tree if n.tag not in (tag('p:nvGrpSpPr'), tag('p:grpSpPr'), tag('p:extLst'))]
        next_id = max(int(n.get('id')) for n in tree.iter(tag('p:cNvPr'))) + 1
        group = el('p:grpSp', children=[el('p:nvGrpSpPr', children=[el('p:cNvPr', {'id': next_id, 'name': 'SVG viewport'}), el('p:cNvGrpSpPr'), el('p:nvPr')]), el('p:grpSpPr', children=[el('a:xfrm', children=[el('a:off', {'x': 0, 'y': 0}), el('a:ext', {'cx': prs.slide_width, 'cy': prs.slide_height}), el('a:chOff', {'x': 0, 'y': 0}), el('a:chExt', {'cx': round(width * EMU_PER_PX), 'cy': round(height * EMU_PER_PX)})])])])
        for node in content:
            tree.remove(node)
            group.append(node)
        tree.append(group)
        drawing = etree.tostring(page, encoding='unicode')
    slide_name = str(slide.part.partname).lstrip("/")
    contents[slide_name] = drawing.encode("utf-8")
    rels_name = str(Path(slide_name).parent / "_rels" / (Path(slide_name).name + ".rels"))
    rels = etree.fromstring(contents[rels_name], parser)
    for rel in relationships:
        etree.SubElement(rels, f"{{{_RELS_NS}}}Relationship", Id=rel["id"],
                         Type=rel["type"], Target=rel["target"])
    contents[rels_name] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8")
    types = etree.fromstring(contents["[Content_Types].xml"], parser)
    extensions = {node.get("Extension") for node in types if node.tag == f"{{{_CT_NS}}}Default"}
    for name, data in media.items():
        contents["ppt/media/" + name] = data
        extension = Path(name).suffix.lstrip(".").lower()
        if extension not in _MEDIA_TYPES:
            raise ValueError(f"Unsupported SVG image format: {extension}")
        if extension not in extensions:
            etree.SubElement(types, f"{{{_CT_NS}}}Default", Extension=extension,
                             ContentType=_MEDIA_TYPES[extension])
            extensions.add(extension)
    contents["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8")
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in contents.items():
            package.writestr(name, data)
    if conversion_trace_path:
        Path(conversion_trace_path).write_text(json.dumps(traces, ensure_ascii=False, indent=2), encoding="utf-8")
    return True
