"""PowerPoint package inspection, validation, rendering and page-part copying."""

from __future__ import annotations
from copy import deepcopy
import hashlib
from pathlib import Path, PurePosixPath
import posixpath
import shutil
import subprocess
import tempfile
from typing import Any
from urllib.parse import unquote
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, _Relationship
from pptx.opc.packuri import PackURI

EMU_PER_PT = 12700
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


class PptxError(ValueError):
    """A package or rendering operation cannot complete."""


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _source(path: str | Path) -> Path:
    source = Path(path).expanduser().resolve(strict=True)
    validate_pptx_package(source)
    return source


def _walk_shapes(shapes: Any):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)


def _frame(shape: Any) -> dict[str, float | None]:
    return {
        key: None
        if getattr(shape, key) is None
        else round(float(getattr(shape, key)) / EMU_PER_PT, 4)
        for key in ("left", "top", "width", "height")
    }


def _text_inventory(frame: Any) -> list[dict[str, Any]]:
    return [
        {
            "text": paragraph.text,
            "alignment": str(paragraph.alignment),
            "runs": [
                {
                    "text": run.text,
                    "font_family": run.font.name,
                    "font_size": float(run.font.size) / EMU_PER_PT
                    if run.font.size
                    else None,
                    "bold": run.font.bold,
                    "italic": run.font.italic,
                }
                for run in paragraph.runs
            ],
        }
        for paragraph in frame.paragraphs
    ]


def inspect_pptx(path: str | Path) -> dict[str, Any]:
    """Return slide-local shape IDs, editable content, layouts and geometry."""
    source = _source(path)
    from .document.identity import read_identity

    prs = Presentation(str(source))
    slides = []
    for number, slide in enumerate(prs.slides, 1):
        shapes = []
        parents = {
            child.shape_id: group.shape_id
            for group in _walk_shapes(slide.shapes)
            if group.shape_type == MSO_SHAPE_TYPE.GROUP
            for child in group.shapes
        }
        for shape in _walk_shapes(slide.shapes):
            item = {
                "shape_id": shape.shape_id,
                "name": shape.name,
                "id": read_identity(
                    shape._element, f"slide-{number}-object-{shape.shape_id}"
                ),
                "type": str(shape.shape_type),
                "frame": _frame(shape),
                "rotation": shape.rotation,
                "is_placeholder": shape.is_placeholder,
                "parent_shape_id": parents.get(shape.shape_id),
                "coordinate_space": "group" if shape.shape_id in parents else "slide",
            }
            if shape.is_placeholder:
                item["placeholder"] = {
                    "idx": shape.placeholder_format.idx,
                    "type": str(shape.placeholder_format.type),
                }
            if shape.has_text_frame:
                item.update(
                    kind="text",
                    text=shape.text,
                    paragraphs=_text_inventory(shape.text_frame),
                )
            elif shape.has_table:
                item.update(
                    kind="table",
                    values=[
                        [cell.text for cell in row.cells] for row in shape.table.rows
                    ],
                )
            elif shape.has_chart:
                item.update(
                    kind="chart",
                    chart_type=str(shape.chart.chart_type),
                    series=[
                        {"name": series.name, "values": list(series.values)}
                        for series in shape.chart.series
                    ],
                )
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                item.update(
                    kind="image",
                    content_type=shape.image.content_type,
                    image_size=list(shape.image.size),
                    crop={
                        key: getattr(shape, f"crop_{key}")
                        for key in ("left", "top", "right", "bottom")
                    },
                )
            else:
                item["kind"] = (
                    "group" if shape.shape_type == MSO_SHAPE_TYPE.GROUP else "shape"
                )
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    transform = shape._element.grpSpPr.xfrm
                    if transform is not None:
                        item["child_frame"] = {
                            "left": transform.chOff.x / EMU_PER_PT,
                            "top": transform.chOff.y / EMU_PER_PT,
                            "width": transform.chExt.cx / EMU_PER_PT,
                            "height": transform.chExt.cy / EMU_PER_PT,
                        }
            shapes.append(item)
        # python-pptx omits drawing entries inside markup compatibility choices.
        # Inventory them explicitly rather than hiding preserved Office Math.
        known_ids = {item["shape_id"] for item in shapes}
        containers = {
            qn("p:sp"),
            qn("p:pic"),
            qn("p:grpSp"),
            qn("p:graphicFrame"),
            qn("p:cxnSp"),
        }
        for properties in slide._element.spTree.iter(qn("p:cNvPr")):
            shape_id = int(properties.get("id"))
            if shape_id in known_ids:
                continue
            owner = properties.getparent()
            while owner is not None and owner.tag not in containers:
                owner = owner.getparent()
            if owner is None:
                continue
            transform = owner.find("./" + qn("p:spPr") + "/" + qn("a:xfrm"))
            frame = None
            if transform is not None:
                off, ext = transform.find(qn("a:off")), transform.find(qn("a:ext"))
                if off is not None and ext is not None:
                    frame = {
                        "left": int(off.get("x")) / EMU_PER_PT,
                        "top": int(off.get("y")) / EMU_PER_PT,
                        "width": int(ext.get("cx")) / EMU_PER_PT,
                        "height": int(ext.get("cy")) / EMU_PER_PT,
                    }
            parent_group = next(
                (node for node in owner.iterancestors() if node.tag == qn("p:grpSp")),
                None,
            )
            parent_id = (
                None
                if parent_group is None
                else int(
                    parent_group.find(
                        "./" + qn("p:nvGrpSpPr") + "/" + qn("p:cNvPr")
                    ).get("id")
                )
            )
            math_namespace = (
                "http://schemas.openxmlformats.org/officeDocument/2006/math"
            )
            is_math = any(
                node.tag == f"{{{math_namespace}}}oMath" for node in owner.iter()
            )
            shapes.append(
                {
                    "shape_id": shape_id,
                    "name": properties.get("name", ""),
                    "kind": "office_math" if is_math else "preserved_drawing",
                    "frame": frame,
                    "parent_shape_id": parent_id,
                    "coordinate_space": "group" if parent_id is not None else "slide",
                    "editable_via_text_api": False,
                }
            )
            known_ids.add(shape_id)
        slides.append(
            {
                "slide": number,
                "id": read_identity(slide._element.cSld, f"slide-{number}"),
                "part": str(slide.part.partname),
                "layout_id": str(slide.slide_layout.part.partname),
                "shapes": shapes,
            }
        )
    layouts = [
        {
            "layout_id": str(layout.part.partname),
            "name": layout.name,
            "master_id": str(master.part.partname),
            "placeholders": [
                {
                    "idx": shape.placeholder_format.idx,
                    "name": shape.name,
                    "frame": _frame(shape),
                    "type": str(shape.placeholder_format.type),
                }
                for shape in layout.placeholders
            ],
        }
        for master in prs.slide_masters
        for layout in master.slide_layouts
    ]
    return {
        "schema": "presentation.inventory.v1",
        "source_pptx": str(source),
        "source_sha256": _sha256(source),
        "units": "pt",
        "slide_size": {
            "width": prs.slide_width / EMU_PER_PT,
            "height": prs.slide_height / EMU_PER_PT,
        },
        "layouts": layouts,
        "slides": slides,
    }


def validate_pptx_package(path: str | Path) -> dict[str, Any]:
    """Check actual package XML and internal relationship destinations."""
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        if len(names) != len(set(names)):
            raise PptxError("PPTX contains duplicate package entries")
        required = {"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"}
        if not required.issubset(names):
            raise PptxError("File is not a PPTX presentation package")
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        xml_files = {
            name: etree.fromstring(package.read(name), parser)
            for name in names
            if name.endswith((".xml", ".rels"))
        }
        relationship_ids = {}
        relationship_targets = {}
        for name, root in xml_files.items():
            if not name.endswith(".rels"):
                continue
            base = posixpath.dirname(posixpath.dirname(name))
            ids = [rel.get("Id") for rel in root]
            if len(ids) != len(set(ids)) or None in ids:
                raise PptxError(f"Invalid relationship IDs in {name}")
            relationship_ids[name] = set(ids)
            relationship_targets[name] = {}
            for rel in root:
                if rel.get("TargetMode") == "External":
                    continue
                target = unquote(rel.get("Target", "").split("#", 1)[0])
                resolved = (
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.normpath(posixpath.join(base, target))
                )
                if resolved not in names:
                    raise PptxError(f"Broken PPTX relationship: {name} -> {target}")
                relationship_targets[name][rel.get("Id")] = {
                    "path": resolved,
                    "type": rel.get("Type"),
                }
        for name, root in xml_files.items():
            if name.endswith(".rels"):
                continue
            path = PurePosixPath(name)
            rels_name = str(path.parent / "_rels" / (path.name + ".rels"))
            ids = relationship_ids.get(rels_name, set())
            for node in root.iter():
                for attr, value in node.attrib.items():
                    if attr.startswith("{" + _R_NS + "}") and value not in ids:
                        raise PptxError(
                            f"Unknown resource relationship in {name}: {value}"
                        )
        presentation = xml_files["ppt/presentation.xml"]
        refs = relationship_targets["ppt/_rels/presentation.xml.rels"]
        masters = {
            rel["path"] for rel in refs.values() if rel["type"] == RT.NOTES_MASTER
        }
        for node in presentation.findall(qn("p:sldIdLst") + "/" + qn("p:sldId")):
            slide = PurePosixPath(refs[node.get(qn("r:id"))]["path"])
            slide_rels = str(slide.parent / "_rels" / (slide.name + ".rels"))
            for note in relationship_targets.get(slide_rels, {}).values():
                if note["type"] != RT.NOTES_SLIDE:
                    continue
                path = PurePosixPath(note["path"])
                note_rels = str(path.parent / "_rels" / (path.name + ".rels"))
                for master in relationship_targets.get(note_rels, {}).values():
                    if (
                        master["type"] == RT.NOTES_MASTER
                        and master["path"] not in masters
                    ):
                        raise PptxError(
                            f"Notes master {master['path']} is not registered in the presentation; copied notes pages must share the registered notes master"
                        )
    return {"status": "ok", "package_entries": len(names), "xml_parts": len(xml_files)}


def copy_slide(prs: Any, source: Any):
    """Copy a slide, cloning mutable chart/workbook/notes parts with their rIds."""
    target = prs.slides.add_slide(source.slide_layout)
    target.part._element = deepcopy(source._element)
    from .document.identity import renew_identities

    renew_identities(target.part._element)
    # Slide wraps its element separately from SlidePart.
    target._element = target.part._element
    # add_slide() populates placeholders through a cached SlideShapes facade.
    # Rebind that facade to the copied shape tree before further edits.
    target.__dict__.pop("shapes", None)
    target.__dict__.pop("placeholders", None)
    target.part.rels._rels.clear()
    mapping = {source.part: target.part}
    reserved = {str(part.partname) for part in prs.part.package.iter_parts()}
    shared = {
        RT.SLIDE_LAYOUT,
        RT.SLIDE_MASTER,
        RT.NOTES_MASTER,
        RT.THEME,
        RT.IMAGE,
        RT.SLIDE,
    }

    def clone(part: Part):
        if part in mapping:
            return mapping[part]
        path = PurePosixPath(str(part.partname))
        number = 1
        candidate = str(path.with_name(f"{path.stem}_pptx{number}{path.suffix}"))
        while candidate in reserved:
            number += 1
            candidate = str(path.with_name(f"{path.stem}_pptx{number}{path.suffix}"))
        reserved.add(candidate)
        copied = type(part).load(
            PackURI(candidate), part.content_type, prs.part.package, part.blob
        )
        mapping[part] = copied
        relationships(part, copied)
        return copied

    def relationships(original: Part, copied: Part):
        for rel in original.rels.values():
            if rel.is_external:
                destination = rel.target_ref
            elif rel.target_part in mapping:
                destination = mapping[rel.target_part]
            elif rel.reltype in shared:
                destination = rel.target_part
            else:
                destination = clone(rel.target_part)
            copied.rels._rels[rel.rId] = _Relationship(
                copied.partname.baseURI,
                rel.rId,
                rel.reltype,
                "External" if rel.is_external else "Internal",
                destination,
            )

    relationships(source.part, target.part)
    return target


def render_pptx(
    source_path: str | Path,
    output_dir: str | Path,
    *,
    renderer: str | None = None,
    rasterizer: str | None = None,
    dpi: int = 96,
    embedded_fonts: str = "ignore",
) -> dict[str, Any]:
    """Render the actual PPTX through LibreOffice and Poppler, with an isolated profile."""
    source = _source(source_path)
    if not 24 <= dpi <= 600:
        raise ValueError("Preview DPI must be between 24 and 600")
    soffice = renderer or shutil.which("soffice")
    pdftoppm = rasterizer or shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        raise PptxError(
            "PPTX preview requires LibreOffice (soffice) and Poppler (pdftoppm)"
        )
    output = Path(output_dir).expanduser().resolve(strict=False)
    if output.exists() and any(output.iterdir()):
        raise PptxError("Choose an empty preview directory to avoid stale slide images")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".pptx-pptx-render-", dir=output.parent
    ) as temp_name:
        temp = Path(temp_name)
        from .document.render_env import office_environment, office_input

        render_source, input_report = office_input(source, temp, embedded_fonts)
        environment, fonts = office_environment(Presentation(str(render_source)), temp)
        fonts.update(input_report)
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={(temp / 'profile').as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temp),
                str(render_source),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=180,
            env=environment,
        )
        pdf = temp / f"{source.stem}.pdf"
        if not pdf.is_file():
            raise PptxError("LibreOffice returned without producing a PDF")
        subprocess.run(
            [pdftoppm, "-png", "-r", str(dpi), str(pdf), str(temp / "slide")],
            check=True,
            capture_output=True,
            text=True,
            timeout=180,
        )
        previews = sorted(temp.glob("slide-*.png"))
        if len(previews) != len(Presentation(str(source)).slides):
            raise PptxError("Rendered page count does not match the presentation")
        output.mkdir(exist_ok=True)
        shutil.copy2(pdf, output / pdf.name)
        for preview in previews:
            shutil.copy2(preview, output / preview.name)
    return {
        "schema": "pptx.preview.v1",
        "source_pptx": str(source),
        "fonts": fonts,
        "pdf": str(output / pdf.name),
        "slides": [str(output / preview.name) for preview in previews],
    }
