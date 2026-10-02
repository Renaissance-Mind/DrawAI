"""Native rich text, geometry, effects and 3-D bindings."""

from __future__ import annotations
import uuid
from io import BytesIO
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
from .xml import (
    EMU,
    P_ORDER,
    RUN_ORDER,
    asset_data,
    clear_choice,
    color,
    el,
    ensure,
    ordered,
    paint,
    replace,
    set_paint,
    spacing,
    stroke,
    tag,
)

RUN_ATTRS = {
    "font.size": ("sz", 100),
    "font.bold": ("b", None),
    "font.italic": ("i", None),
    "underline.type": ("u", None),
    "strike": ("strike", None),
    "capitalization": ("cap", None),
    "letter_spacing": ("spc", 100),
    "kerning_threshold": ("kern", 100),
    "baseline_shift": ("baseline", 100000),
    "normalize_height": ("normalizeH", None),
    "kumimoji": ("kumimoji", None),
    "language": ("lang", None),
    "alternative_language": ("altLang", None),
    "bookmark": ("bmk", None),
}
P_ATTRS = {
    "alignment": ("algn", None),
    "level": ("lvl", None),
    "left_margin": ("marL", EMU),
    "right_margin": ("marR", EMU),
    "first_line_indent": ("indent", EMU),
    "default_tab_size": ("defTabSz", EMU),
    "rtl": ("rtl", None),
    "font_alignment": ("fontAlgn", None),
    "east_asian_line_break": ("eaLnBrk", None),
    "latin_line_break": ("latinLnBrk", None),
    "hanging_punctuation": ("hangingPunct", None),
}
BODY_ATTRS = {
    "vertical_anchor": ("anchor", None),
    "anchor_center": ("anchorCtr", None),
    "wrap": ("wrap", None),
    "rotation": ("rot", 60000),
    "vertical_mode": ("vert", None),
    "columns": ("numCol", None),
    "column_spacing": ("spcCol", EMU),
    "columns_rtl": ("rtlCol", None),
    "upright": ("upright", None),
    "vertical_overflow": ("vertOverflow", None),
    "horizontal_overflow": ("horzOverflow", None),
    "first_last_paragraph_spacing": ("spcFirstLastPara", None),
    "compatible_line_spacing": ("compatLnSpc", None),
    "force_antialias": ("forceAA", None),
}


def attrs(node, props, prefix, bindings):
    used = set()
    for key, (name, scale) in bindings.items():
        full = prefix + key
        if full in props:
            value = props[full]
            if scale is not None:
                value = round(value * scale)
            node.set(
                name, "1" if value is True else "0" if value is False else str(value)
            )
            used.add(full)
    return used


def hyperlink(node, name, value, part):
    attr = {
        k: value[k]
        for k in (
            "tooltip",
            "tgtFrame",
            "history",
            "highlightClick",
            "endSnd",
            "action",
        )
        if k in value
    }
    if "url" in value:
        attr["r:id"] = part.relate_to(value["url"], RT.HYPERLINK, is_external=True)
    replace(node, name, el(name, attr))


def run_style(node, props, part, assets, root):
    used = attrs(node, props, "text.run.", RUN_ATTRS)
    for key, name in [
        ("latin", "a:latin"),
        ("east_asian", "a:ea"),
        ("complex_script", "a:cs"),
        ("symbol", "a:sym"),
    ]:
        full = "text.run.font." + key
        if full in props:
            value = props[full]
            if isinstance(value, str):
                value = {"typeface": value}
            value = dict(value)
            for field, native in [
                ("text.run.font.charset", "charset"),
                ("text.run.font.pitch_family", "pitchFamily"),
            ]:
                if field in props:
                    value[native] = props[field]
                    used.add(field)
            replace(node, name, el(name, value))
            used.add(full)
    for full, name, builder in [
        ("text.run.fill", "a:solidFill", paint),
        ("text.run.outline", "a:ln", stroke),
    ]:
        if full in props:
            if full.endswith("fill"):
                set_paint(node, props[full], part, assets, root)
            else:
                replace(node, name, builder(props[full], part, assets, root))
            used.add(full)
    if "text.run.highlight" in props:
        replace(
            node,
            "a:highlight",
            el("a:highlight", children=[color(props["text.run.highlight"])]),
        )
        used.add("text.run.highlight")
    for full, name in [
        ("text.run.underline.fill", "a:uFill"),
        ("text.run.underline.stroke", "a:uLn"),
    ]:
        if full in props:
            value = props[full]
            child = (
                paint(value, part, assets, root)
                if name == "a:uFill"
                else stroke(value, part, assets, root)
            )
            replace(
                node,
                name,
                el(name, children=[child])
                if name == "a:uFill"
                else el(name, dict(child.attrib), list(child)),
            )
            used.add(full)
    for full, name in [
        ("text.run.click_link", "a:hlinkClick"),
        ("text.run.hover_link", "a:hlinkMouseOver"),
    ]:
        if full in props:
            hyperlink(node, name, props[full], part)
            used.add(full)
    if "text.run.rtl" in props:
        replace(node, "a:rtl", el("a:rtl", {"val": props["text.run.rtl"]}))
        used.add("text.run.rtl")
    if "text.run.proofing" in props:
        for key, value in props["text.run.proofing"].items():
            if key not in ("noProof", "dirty", "err", "smtClean", "smtId"):
                raise ValueError(f"Invalid proofing property {key}")
            node.set(
                key, "1" if value is True else "0" if value is False else str(value)
            )
        used.add("text.run.proofing")
    if "text.run.effects" in props:
        effect_list(node, props["text.run.effects"])
        used.add("text.run.effects")
    ordered(node, RUN_ORDER)
    missing = props.keys() - used
    if missing:
        raise ValueError(f"Unsupported nested run settings: {sorted(missing)}")
    return used


def paragraph(node, value, part, assets, root, default=None):
    props = value.get("properties", {})
    if value.get("native_ppr") and "baseline_properties" in value:
        props = {
            k: v for k, v in props.items() if value["baseline_properties"].get(k) != v
        }
    ppr = (
        etree.fromstring(
            value["native_ppr"].encode(),
            etree.XMLParser(resolve_entities=False, no_network=True),
        )
        if value.get("native_ppr")
        else el("a:pPr")
    )
    node.append(ppr)
    used = attrs(ppr, props, "text.paragraph.", P_ATTRS)
    for key, name in [
        ("line_spacing", "lnSpc"),
        ("space_before", "spcBef"),
        ("space_after", "spcAft"),
    ]:
        full = "text.paragraph." + key
        if full in props:
            replace(ppr, "a:" + name, spacing(name, props[full]))
            used.add(full)
    if "text.paragraph.tab_stops" in props:
        ppr.append(
            el(
                "a:tabLst",
                children=[
                    el(
                        "a:tab",
                        {
                            "pos": round(t["position"] * EMU),
                            "algn": t.get("alignment", "l"),
                        },
                    )
                    for t in props["text.paragraph.tab_stops"]
                ],
            )
        )
        used.add("text.paragraph.tab_stops")
    run_default = dict(default or {})
    run_default.update(props.get("text.paragraph.default_run_style", {}))
    if run_default:
        d = ensure(ppr, "a:defRPr")
        used.update(run_style(d, run_default, part, assets, root))
    if "text.paragraph.default_run_style" in props:
        used.add("text.paragraph.default_run_style")
    bullet = "text.paragraph.bullet."
    if bullet + "color" in props:
        ppr.append(el("a:buClr", children=[color(props[bullet + "color"])]))
        used.add(bullet + "color")
    if bullet + "font" in props:
        ppr.append(el("a:buFont", {"typeface": props[bullet + "font"]}))
        used.add(bullet + "font")
    if bullet + "size" in props:
        v = props[bullet + "size"]
        ppr.append(
            el(
                "a:buSzPts" if v["mode"] == "points" else "a:buSzPct",
                {"val": round(v["value"] * (100 if v["mode"] == "points" else 100000))},
            )
        )
        used.add(bullet + "size")
    if bullet + "kind" in props:
        k = props[bullet + "kind"]
        used.add(bullet + "kind")
        if k == "none":
            ppr.append(el("a:buNone"))
        elif k == "character":
            ppr.append(el("a:buChar", {"char": props[bullet + "character"]}))
            used.add(bullet + "character")
        elif k == "auto_number":
            v = props[bullet + "numbering"]
            ppr.append(
                el("a:buAutoNum", {"type": v["type"], "startAt": v.get("start", 1)})
            )
            used.add(bullet + "numbering")
        elif k == "picture":
            _, rid = part.get_or_add_image_part(
                BytesIO(asset_data(assets[props[bullet + "image_ref"]], root))
            )
            ppr.append(el("a:buBlip", children=[el("a:blip", {"r:embed": rid})]))
            used.add(bullet + "image_ref")
    ordered(ppr, P_ORDER)
    if "text.paragraph.content" in props:
        used.add("text.paragraph.content")
    for content in value.get("content", props.get("text.paragraph.content", [])):
        kind = content.get("kind", "run")
        rp = content.get("properties", {})
        if content.get("native_rpr") and "baseline_properties" in content:
            rp = {
                k: v
                for k, v in rp.items()
                if content["baseline_properties"].get(k) != v
            }
        if kind == "math":
            # Native Office Math can be supplied without flattening into a picture.
            xml = etree.fromstring(
                content["omml"].encode(),
                etree.XMLParser(resolve_entities=False, no_network=True),
            )
            if xml.tag == tag("m:oMath"):
                xml = el("m:oMathPara", children=[xml])
            if xml.tag == tag("m:oMathPara"):
                xml = el("a14:m", children=[xml])
            if xml.tag != tag("a14:m"):
                raise ValueError("Math content must be an Office Math tree")
            node.append(xml)
            continue
        if kind == "run":
            child = el("a:r")
        elif kind == "break":
            child = el("a:br")
        elif kind == "field":
            child = el(
                "a:fld",
                {
                    "id": content.get("id", "{" + str(uuid.uuid4()).upper() + "}"),
                    "type": content["type"],
                },
            )
        else:
            raise ValueError(f"Unsupported paragraph content {kind}")
        rpr = (
            etree.fromstring(
                content["native_rpr"].encode(),
                etree.XMLParser(resolve_entities=False, no_network=True),
            )
            if content.get("native_rpr")
            else el("a:rPr")
        )
        child.append(rpr)
        used.update(run_style(rpr, rp, part, assets, root))
        if kind != "break":
            t = el("a:t")
            t.text = content.get("text", "")
            child.append(t)
        node.append(child)
    if "text.paragraph.end_run_style" in props:
        end = el("a:endParaRPr")
        used.update(
            run_style(end, props["text.paragraph.end_run_style"], part, assets, root)
        )
        node.append(end)
        used.add("text.paragraph.end_run_style")
    for xml in value.get("native_tail", []):
        node.append(
            etree.fromstring(
                xml.encode(), etree.XMLParser(resolve_entities=False, no_network=True)
            )
        )
    missing = props.keys() - used
    if missing:
        raise ValueError(f"Unsupported nested paragraph settings: {sorted(missing)}")
    return used, props


def text_body(body, props, part, assets, root):
    used = set()
    bpr = ensure(body, "a:bodyPr")
    used.update(attrs(bpr, props, "text.body.", BODY_ATTRS))
    if "text.body.insets" in props:
        for key, value in props["text.body.insets"].items():
            bpr.set(
                {"left": "lIns", "right": "rIns", "top": "tIns", "bottom": "bIns"}[key],
                str(round(value * EMU)),
            )
        used.add("text.body.insets")
    if "text.body.autofit" in props:
        clear_choice(bpr, ["a:noAutofit", "a:normAutofit", "a:spAutoFit"])
        value = props["text.body.autofit"]
        if isinstance(value, dict):
            value = value["mode"]
        value = {"shrink_text": "shrink", "grow_shape": "resize"}.get(value, value)
        name = {
            "none": "a:noAutofit",
            "shrink": "a:normAutofit",
            "resize": "a:spAutoFit",
        }[value]
        a = {}
        if value == "shrink" and "text.body.autofit_scale" in props:
            v = props["text.body.autofit_scale"]
            a = {
                "fontScale": round(v.get("font", 1) * 100000),
                "lnSpcReduction": round(v.get("line_spacing_reduction", 0) * 100000),
            }
            used.add("text.body.autofit_scale")
        bpr.append(el(name, a))
        used.add("text.body.autofit")
    if "text.body.wordart_warp" in props:
        v = props["text.body.wordart_warp"]
        replace(
            bpr,
            "a:prstTxWarp",
            el(
                "a:prstTxWarp",
                {"prst": v["preset"]},
                [guide_list("avLst", v.get("adjustments", []))],
            ),
        )
        used.add("text.body.wordart_warp")
    ensure(body, "a:lstStyle")
    default = {k: v for k, v in props.items() if k.startswith("text.run.")}
    if "text.body.paragraphs" in props:
        for p in list(body):
            if p.tag == tag("a:p"):
                body.remove(p)
        for p in props["text.body.paragraphs"] or [{"content": []}]:
            child = el("a:p")
            textused, _ = paragraph(child, p, part, assets, root, default)
            body.append(child)
            used.update(textused)
        used.add("text.body.paragraphs")
        used.update(default)
    elif default:
        for p in body.findall(tag("a:p")):
            ppr = ensure(p, "a:pPr")
            dpr = ensure(ppr, "a:defRPr")
            used.update(run_style(dpr, default, part, assets, root))
            ordered(ppr, P_ORDER)
            for r in p:
                if r.tag in (tag("a:r"), tag("a:br"), tag("a:fld")):
                    used.update(
                        run_style(ensure(r, "a:rPr"), default, part, assets, root)
                    )
    ordered(
        bpr,
        [
            "a:prstTxWarp",
            "a:noAutofit",
            "a:normAutofit",
            "a:spAutoFit",
            "a:scene3d",
            "a:sp3d",
            "a:flatTx",
            "a:extLst",
        ],
    )
    return used


def guide_list(name, values):
    if isinstance(values, dict):
        values = [{"name": k, "formula": v} for k, v in values.items()]
    return el(
        "a:" + name,
        children=[
            el("a:gd", {"name": v["name"], "fmla": v["formula"]}) for v in values
        ],
    )


def geometry(sp, props):
    used = set()
    if "shape.geometry.preset" in props:
        clear_choice(sp, ["a:prstGeom", "a:custGeom"])
        sp.append(
            el(
                "a:prstGeom",
                {"prst": props["shape.geometry.preset"]},
                [guide_list("avLst", props.get("shape.geometry.adjustments", []))],
            )
        )
        used.add("shape.geometry.preset")
        if "shape.geometry.adjustments" in props:
            used.add("shape.geometry.adjustments")
    if "shape.geometry.paths" in props:
        clear_choice(sp, ["a:prstGeom", "a:custGeom"])
        g = el("a:custGeom")
        for key, name in [("adjustments", "avLst"), ("guides", "gdLst")]:
            g.append(guide_list(name, props.get("shape.geometry." + key, [])))
            if "shape.geometry." + key in props:
                used.add("shape.geometry." + key)
        handles = el("a:ahLst")
        for h in props.get("shape.geometry.handles", []):
            kind = h["kind"]
            node = el(
                "a:ahPolar" if kind == "polar" else "a:ahXY",
                h.get("attributes", {}),
                [el("a:pos", {"x": h["x"], "y": h["y"]})],
            )
            handles.append(node)
        g.append(handles)
        if "shape.geometry.handles" in props:
            used.add("shape.geometry.handles")
        g.append(
            el(
                "a:cxnLst",
                children=[
                    el(
                        "a:cxn",
                        {"ang": s["angle"]},
                        [el("a:pos", {"x": s["x"], "y": s["y"]})],
                    )
                    for s in props.get("shape.geometry.connection_sites", [])
                ],
            )
        )
        if "shape.geometry.connection_sites" in props:
            used.add("shape.geometry.connection_sites")
        rect = props.get(
            "shape.geometry.text_rect", {"l": "l", "t": "t", "r": "r", "b": "b"}
        )
        g.append(el("a:rect", rect))
        if "shape.geometry.text_rect" in props:
            used.add("shape.geometry.text_rect")
        paths = el("a:pathLst")
        for path in props["shape.geometry.paths"]:
            pa = {
                k: path[k]
                for k in ("w", "h", "fill", "stroke", "extrusionOk")
                if k in path
            }
            viewport = props.get("shape.geometry.path_viewport", {})
            for key, native in [
                ("path_stroke", "stroke"),
                ("path_fill", "fill"),
                ("extrusion_ok", "extrusionOk"),
            ]:
                full = "shape.geometry." + key
                if full in props:
                    pa[native] = props[full]
                    used.add(full)
            if viewport:
                pa.setdefault("w", viewport["width"])
                pa.setdefault("h", viewport["height"])
                used.add("shape.geometry.path_viewport")
            node = el("a:path", pa)
            for command in path["commands"]:
                op = {
                    "lineTo": "lnTo",
                    "quadBezierTo": "quadBezTo",
                    "cubicBezierTo": "cubicBezTo",
                }.get(command["op"], command["op"])
                if op in ("moveTo", "lnTo"):
                    node.append(
                        el(
                            "a:" + op,
                            children=[
                                el("a:pt", {"x": command["x"], "y": command["y"]})
                            ],
                        )
                    )
                elif op in ("cubicBezTo", "quadBezTo"):
                    node.append(
                        el(
                            "a:" + op,
                            children=[
                                el("a:pt", {"x": p[0], "y": p[1]})
                                for p in command["points"]
                            ],
                        )
                    )
                elif op == "arcTo":
                    node.append(
                        el(
                            "a:arcTo",
                            {k: command[k] for k in ("wR", "hR", "stAng", "swAng")},
                        )
                    )
                elif op == "close":
                    node.append(el("a:close"))
                else:
                    raise ValueError(f"Unknown path command {op}")
            paths.append(node)
        g.append(paths)
        sp.append(g)
        used.add("shape.geometry.paths")
        if props.get("shape.geometry.path_fill_rule") == "nonzero":
            used.add("shape.geometry.path_fill_rule")
    if (
        "shape.geometry.adjustments" in props
        and "shape.geometry.preset" not in props
        and "shape.geometry.paths" not in props
    ):
        g = sp.find(tag("a:prstGeom"))
        if g is None:
            raise ValueError("Adjustments need a preset or custom geometry")
        replace(g, "a:avLst", guide_list("avLst", props["shape.geometry.adjustments"]))
        used.add("shape.geometry.adjustments")
    return used


def effect_list(parent, effects):
    node = el("a:effectLst")
    for value in effects:
        name = value["kind"]
        a = dict(value.get("attributes", {}))
        for key, mul in [
            ("blur", EMU),
            ("distance", EMU),
            ("direction", 60000),
            ("radius", EMU),
            ("scale_x", 100000),
            ("scale_y", 100000),
            ("skew_x", 60000),
            ("skew_y", 60000),
        ]:
            if key in value:
                a[
                    {
                        "blur": "blurRad",
                        "distance": "dist",
                        "direction": "dir",
                        "radius": "rad",
                        "scale_x": "sx",
                        "scale_y": "sy",
                        "skew_x": "kx",
                        "skew_y": "ky",
                    }[key]
                ] = round(value[key] * mul)
        child = el("a:" + name, a)
        if "color" in value:
            child.append(color(value["color"]))
        if "colors" in value:
            child.extend(color(c) for c in value["colors"])
        if "xml" in value:
            child.extend(etree.fromstring(value["xml"].encode()))
        node.append(child)
    replace(parent, "a:effectLst", node)


def shape_effects(sp, props):
    mapping = {
        "outer_shadow": "outerShdw",
        "inner_shadow": "innerShdw",
        "preset_shadow": "prstShdw",
        "glow": "glow",
        "soft_edge": "softEdge",
        "reflection": "reflection",
        "blur": "blur",
    }
    effects = []
    used = set()
    for key, native in mapping.items():
        full = "effects." + key
        if full in props:
            value = props[full]
            value = (
                {"radius": value} if isinstance(value, (int, float)) else dict(value)
            )
            value["kind"] = native
            effects.append(value)
            used.add(full)
    if effects:
        effect_list(sp, effects)
    if "effects.dag" in props:
        replace(
            sp,
            "a:effectDag",
            etree.fromstring(
                props["effects.dag"].encode(),
                etree.XMLParser(resolve_entities=False, no_network=True),
            ),
        )
        used.add("effects.dag")
    return used


def rotation3d(value):
    return el(
        "a:rot",
        {
            k: round(value.get(n, 0) * 60000)
            for k, n in [
                ("lat", "latitude"),
                ("lon", "longitude"),
                ("rev", "revolution"),
            ]
        },
    )


def three_d(sp, props):
    used = set()
    if any(
        k.startswith("three_d.camera.")
        or k.startswith("three_d.light_rig.")
        or k == "three_d.backdrop"
        for k in props
    ):
        scene = el("a:scene3d")
        cam = el(
            "a:camera",
            {"prst": props.get("three_d.camera.preset", "orthographicFront")},
        )
        for k, a, m in [("field_of_view", "fov", 60000), ("zoom", "zoom", 100000)]:
            if "three_d.camera." + k in props:
                cam.set(a, str(round(props["three_d.camera." + k] * m)))
                used.add("three_d.camera." + k)
        if "three_d.camera.rotation" in props:
            cam.append(rotation3d(props["three_d.camera.rotation"]))
            used.add("three_d.camera.rotation")
        scene.append(cam)
        rig = el(
            "a:lightRig",
            {
                "rig": props.get("three_d.light_rig.preset", "threePt"),
                "dir": props.get("three_d.light_rig.direction", "t"),
            },
        )
        if "three_d.light_rig.rotation" in props:
            rig.append(rotation3d(props["three_d.light_rig.rotation"]))
            used.add("three_d.light_rig.rotation")
        scene.append(rig)
        if "three_d.backdrop" in props:
            v = props["three_d.backdrop"]
            scene.append(
                el(
                    "a:backdrop",
                    children=[
                        el(
                            "a:anchor",
                            {k: round(x * EMU) for k, x in v["anchor"].items()},
                        ),
                        el("a:norm", v["normal"]),
                        el("a:up", v["up"]),
                    ],
                )
            )
            used.add("three_d.backdrop")
        replace(sp, "a:scene3d", scene)
        used.update(
            k
            for k in (
                "three_d.camera.preset",
                "three_d.light_rig.preset",
                "three_d.light_rig.direction",
            )
            if k in props
        )
    bindings = {
        "extrusion_depth": ("extrusionH", EMU),
        "contour_width": ("contourW", EMU),
        "z_position": ("z", EMU),
        "material": ("prstMaterial", None),
    }
    if any(k.startswith("three_d.") and k not in used for k in props):
        node = el("a:sp3d")
        used.update(attrs(node, props, "three_d.", bindings))
        for key, name in [("bevel.top", "a:bevelT"), ("bevel.bottom", "a:bevelB")]:
            if "three_d." + key in props:
                v = props["three_d." + key]
                node.append(
                    el(
                        name,
                        {
                            "w": round(v.get("width", 6) * EMU),
                            "h": round(v.get("height", 6) * EMU),
                            "prst": v.get("preset", "circle"),
                        },
                    )
                )
                used.add("three_d." + key)
        for key, name in [
            ("extrusion_color", "a:extrusionClr"),
            ("contour_color", "a:contourClr"),
        ]:
            if "three_d." + key in props:
                node.append(el(name, children=[color(props["three_d." + key])]))
                used.add("three_d." + key)
        replace(sp, "a:sp3d", node)
    return used
