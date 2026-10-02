"""Lossless template envelope plus normalized, editable object properties."""

from __future__ import annotations
import base64
from copy import deepcopy
import hashlib
import json
import zlib
from io import BytesIO
from pathlib import Path
import tempfile
from pptx import Presentation
from .model import Document, Page, Object
from lxml import etree
from .xml import EMU, tag
from .rich import RUN_ATTRS, P_ATTRS, BODY_ATTRS
from .identity import read_identity


def read_attrs(node, prefix, bindings):
    result = {}
    if node is None:
        return result
    booleans = {
        "b",
        "i",
        "normalizeH",
        "kumimoji",
        "rtl",
        "eaLnBrk",
        "latinLnBrk",
        "hangingPunct",
        "anchorCtr",
        "rtlCol",
        "upright",
        "spcFirstLastPara",
        "compatLnSpc",
        "forceAA",
    }
    for key, (name, scale) in bindings.items():
        if name not in node.attrib:
            continue
        v = node.get(name)
        if scale:
            v = int(v) / scale
        elif name in booleans:
            v = v in ("1", "true")
        elif name in ("lvl", "numCol"):
            v = int(v)
        result[prefix + key] = v
    return result


def read_color(node):
    if node is None:
        return None
    names = {
        "srgbClr": "rgb",
        "scrgbClr": "scrgb",
        "hslClr": "hsl",
        "schemeClr": "scheme",
        "sysClr": "system",
        "prstClr": "preset",
    }
    name = etree.QName(node).localname
    kind = names[name]
    if kind == "scrgb":
        v = [int(node.get(k)) / 100000 for k in ("r", "g", "b")]
    elif kind == "hsl":
        v = [
            int(node.get("hue")) / 60000,
            int(node.get("sat")) / 100000,
            int(node.get("lum")) / 100000,
        ]
    else:
        v = node.get("val")
    result = {"kind": kind, "channels": v}
    if kind == "system":
        result["fallback"] = node.get("lastClr", "000000")
    transforms = []
    for t in node:
        key = etree.QName(t).localname
        if key == "alpha":
            result["alpha"] = int(t.get("val")) / 100000
        else:
            transforms.append(
                {
                    "kind": key,
                    **(
                        {
                            "value": int(t.get("val"))
                            / (60000 if key in ("hue", "hueOff") else 100000)
                        }
                        if "val" in t.attrib
                        else {}
                    ),
                }
            )
    if transforms:
        result["transforms"] = transforms
    return result


def read_paint(parent):
    for node in parent:
        if node.tag == tag("a:noFill"):
            return {"kind": "none"}
        if node.tag == tag("a:grpFill"):
            return {"kind": "group_inherit"}
        if node.tag == tag("a:solidFill"):
            return {"kind": "solid", "color": read_color(node[0])}
        if node.tag == tag("a:gradFill"):
            v = {
                "kind": "gradient",
                "stops": [
                    {"position": int(s.get("pos")) / 100000, "color": read_color(s[0])}
                    for s in node.find(tag("a:gsLst"))
                ],
            }
            lin = node.find(tag("a:lin"))
            path = node.find(tag("a:path"))
            if lin is not None:
                v.update(
                    gradient_kind="linear",
                    angle=int(lin.get("ang", "0")) / 60000,
                    scaled=lin.get("scaled", "1") == "1",
                )
            elif path is not None:
                v.update(gradient_kind=path.get("path", "circle"))
            v["rotate_with_shape"] = node.get("rotWithShape", "1") == "1"
            if node.get("flip"):
                v["flip"] = node.get("flip")
            return v
        if node.tag == tag("a:pattFill"):
            return {
                "kind": "pattern",
                "preset": node.get("prst"),
                "foreground": read_color(node.find(tag("a:fgClr"))[0]),
                "background": read_color(node.find(tag("a:bgClr"))[0]),
            }
    return None


def read_run(node):
    result = read_attrs(node, "text.run.", RUN_ATTRS)
    if node is None:
        return result
    for key, name in [
        ("latin", "a:latin"),
        ("east_asian", "a:ea"),
        ("complex_script", "a:cs"),
        ("symbol", "a:sym"),
    ]:
        child = node.find(tag(name))
        if child is not None:
            result["text.run.font." + key] = dict(child.attrib)
    p = read_paint(node)
    if p:
        result["text.run.fill"] = p
    return result


def read_text(body):
    result = read_attrs(body.find(tag("a:bodyPr")), "text.body.", BODY_ATTRS)
    paras = []
    for p in body.findall(tag("a:p")):
        pp = p.find(tag("a:pPr"))
        props = read_attrs(pp, "text.paragraph.", P_ATTRS)
        if pp is not None:
            dr = pp.find(tag("a:defRPr"))
            if dr is not None:
                props["text.paragraph.default_run_style"] = read_run(dr)
            for key, name in [
                ("line_spacing", "lnSpc"),
                ("space_before", "spcBef"),
                ("space_after", "spcAft"),
            ]:
                space = pp.find(tag("a:" + name))
                if space is not None and len(space):
                    pts = space[0].tag == tag("a:spcPts")
                    props["text.paragraph." + key] = {
                        "mode": "points" if pts else "ratio",
                        "value": int(space[0].get("val")) / (100 if pts else 100000),
                    }
        content = []
        tail = []
        for r in p:
            if r.tag == tag("a14:m"):
                content.append(
                    {"kind": "math", "omml": etree.tostring(r, encoding="unicode")}
                )
                continue
            if r.tag not in (tag("a:r"), tag("a:br"), tag("a:fld"), tag("a:pPr")):
                tail.append(etree.tostring(r, encoding="unicode"))
                continue
            if r.tag not in (tag("a:r"), tag("a:br"), tag("a:fld")):
                continue
            item = {
                "kind": "break"
                if r.tag == tag("a:br")
                else "field"
                if r.tag == tag("a:fld")
                else "run",
                "properties": read_run(r.find(tag("a:rPr"))),
            }
            item["baseline_properties"] = deepcopy(item["properties"])
            rpr = r.find(tag("a:rPr"))
            if rpr is not None:
                item["native_rpr"] = etree.tostring(rpr, encoding="unicode")
            if r.tag != tag("a:br"):
                item["text"] = r.findtext(tag("a:t")) or ""
            if item["kind"] == "field":
                item.update(id=r.get("id"), type=r.get("type", ""))
            content.append(item)
        para = {
            "properties": props,
            "baseline_properties": deepcopy(props),
            "content": content,
        }
        if pp is not None:
            para["native_ppr"] = etree.tostring(pp, encoding="unicode")
        if tail:
            para["native_tail"] = tail
        paras.append(para)
    result["text.body.paragraphs"] = paras
    return result


def read_shape(shape, pageid, assets):
    props = {
        "object.name": shape.name,
        "object.frame": {
            k: float(getattr(shape, k)) / EMU
            for k in ("left", "top", "width", "height")
        },
        "object.rotation": shape.rotation,
    }
    kind = "shape"
    if shape.shape_type == 6:
        kind = "group"
    elif shape.shape_type == 13:
        kind = "picture"
    elif shape.has_table:
        kind = "table"
    elif shape.has_chart:
        kind = "chart"
    elif shape.shape_type == 17:
        kind = "text_box"
    elif shape.shape_type == 9:
        kind = "connector"
    sp = shape._element.find(tag("p:spPr"))
    if sp is not None:
        p = read_paint(sp)
        if p:
            props["paint.kind"] = p["kind"]
            if p["kind"] == "solid":
                props["paint.solid.color"] = p["color"]
            elif p["kind"] == "gradient":
                for key, v in p.items():
                    if key == "kind":
                        continue
                    props[
                        "paint.gradient." + ("kind" if key == "gradient_kind" else key)
                    ] = v
            elif p["kind"] == "pattern":
                for key in ("preset", "foreground", "background"):
                    props["paint.pattern." + key] = p[key]
        g = sp.find(tag("a:prstGeom"))
        if g is not None:
            props["shape.geometry.preset"] = g.get("prst")
    if shape.has_text_frame:
        props.update(read_text(shape.text_frame._txBody))
    if kind == "picture":
        ref = f"{pageid}-image-{shape.shape_id}"
        image = shape.image
        assets[ref] = {
            "data": base64.b64encode(image.blob).decode(),
            "content_type": image.content_type,
        }
        props["picture.asset_ref"] = ref
        props["picture.crop"] = {
            k: float(getattr(shape, "crop_" + v))
            for k, v in [("l", "left"), ("r", "right"), ("t", "top"), ("b", "bottom")]
        }
    if kind == "table":
        table = shape.table
        props["table.grid_columns"] = [c.width / EMU for c in table.columns]
        props["table.rows.height"] = [r.height / EMU for r in table.rows]
        props["table.cells.content"] = [
            [read_text(c.text_frame._txBody) for c in r.cells] for r in table.rows
        ]
    if kind == "chart":
        chart = shape.chart
        plot = chart._chartSpace.find(".//" + tag("c:plotArea"))
        plots = [n for n in plot if etree.QName(n).localname.endswith("Chart")]
        if len(plots) == 1 and etree.QName(plots[0]).localname in (
            "barChart",
            "lineChart",
            "areaChart",
            "pieChart",
            "doughnutChart",
            "radarChart",
        ):
            props["chart.type"] = etree.QName(plots[0]).localname
            props["chart.categories"] = [c.label for c in chart.plots[0].categories]
            props["chart.series"] = [
                {"name": s.name, "values": list(s.values)} for s in chart.series
            ]
    # Chart sources, Office Math, media, SmartArt, unknown extensions remain in the package envelope.
    children = (
        [read_shape(c, pageid, assets) for c in shape.shapes] if kind == "group" else []
    )
    return Object(
        id=read_identity(shape._element, f"{pageid}-object-{shape.shape_id}"),
        kind=kind,
        properties=props,
        baseline=deepcopy(props),
        source_shape_id=shape.shape_id,
        children=children,
        native_xml=etree.tostring(shape._element, encoding="unicode"),
    )


def import_pptx(source, *, embed=True):
    source = Path(source).resolve()
    data = source.read_bytes()
    prs = Presentation(BytesIO(data))
    template = {"sha256": hashlib.sha256(data).hexdigest()}
    if embed:
        template["data"] = base64.b64encode(data).decode()
    else:
        template["path"] = str(source)
    props = {
        "document.page_size": {
            "width": prs.slide_width / EMU,
            "height": prs.slide_height / EMU,
        }
    }
    for key in ("title", "subject", "author", "description", "revision", "keywords"):
        v = getattr(prs.core_properties, "comments" if key == "description" else key)
        props["document.metadata." + key] = (
            v.split(", ")
            if key == "keywords" and v
            else []
            if key == "keywords"
            else str(v)
        )
    assets = {}
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        pageid = read_identity(slide._element.cSld, f"slide-{i}")
        sp = {
            "slide.name": slide._element.cSld.get("name", ""),
            "slide.hidden": slide._element.get("show", "1") == "0",
            "slide.layout_ref": str(slide.slide_layout.part.partname),
        }
        pages.append(
            Page(
                id=pageid,
                source_slide=i,
                properties=sp,
                baseline=deepcopy(sp),
                objects=[read_shape(s, pageid, assets) for s in slide.shapes],
            )
        )
    id = (
        prs.core_properties._element.findtext(
            "{http://purl.org/dc/elements/1.1/}identifier"
        )
        or source.stem
    )
    return Document(
        id=id,
        properties=props,
        baseline=deepcopy(props),
        assets=assets,
        pages=pages,
        template=template,
    )


def import_svg(source, *, padding=0):
    from ..svg_export import SvgToPptCompiler

    source = Path(source).resolve()
    # The converter prepares assets beside its input, so work on a copy.
    with tempfile.TemporaryDirectory() as tmp:
        root = etree.fromstring(
            source.read_bytes(),
            etree.XMLParser(resolve_entities=False, no_network=True),
        )
        for metadata in root.findall("{http://www.w3.org/2000/svg}metadata"):
            if not metadata.text or not metadata.text.lstrip().startswith("{"):
                continue
            value = json.loads(metadata.text)
            if (
                value.get("schema") == "presentation.document.v1"
                and value.get("document_encoding") == "zlib-base64"
            ):
                drawing = deepcopy(root)
                drawing.remove(
                    drawing.findall("{http://www.w3.org/2000/svg}metadata")[0]
                )
                digest = hashlib.sha256(
                    etree.tostring(drawing, method="c14n")
                ).hexdigest()
                if digest != value.get("visual_sha256"):
                    raise ValueError(
                        "SVG drawing changed after export; its embedded document is stale. Edit the document model or remove this metadata before visual re-import."
                    )
                raw = base64.b64decode(value["document_data"], validate=True)
                decoder = zlib.decompressobj()
                raw = decoder.decompress(raw, 64 * 1024 * 1024 + 1)
                if len(raw) > 64 * 1024 * 1024 or not decoder.eof:
                    raise ValueError(
                        "Embedded SVG source exceeds the 64 MiB import limit"
                    )
                return Document.model_validate_json(raw), {
                    "status": "ok",
                    "backend": "embedded_document",
                    "lossless_source_recovery": True,
                }
        expand_svg_uses(root)
        if padding:
            if padding < 0:
                raise ValueError("SVG padding must be nonnegative viewBox units")
            x, y, w, h = map(float, root.get("viewBox", "").replace(",", " ").split())
            root.set(
                "viewBox",
                f"{x - padding} {y - padding} {w + 2 * padding} {h + 2 * padding}",
            )
        for image in root.iter("{http://www.w3.org/2000/svg}image"):
            for attr in ("href", "{http://www.w3.org/1999/xlink}href"):
                value = image.get(attr)
                if value and not value.startswith(("data:", "http:", "https:")):
                    image.set(attr, str((source.parent / value).resolve()))
        path = Path(tmp) / "input.svg"
        path.write_bytes(etree.tostring(root))
        out = Path(tmp) / "import.pptx"
        report = SvgToPptCompiler().compile(
            path, out, check_full_slide_editability=False
        )
        if report["status"] != "ok":
            raise ValueError(f"SVG import failed: {report}")
        doc = import_pptx(out)
        doc.id = source.stem
    return doc, report


def expand_svg_uses(root):
    """Resolve reusable vector glyphs such as dvisvgm's <use> references."""
    ns = "{http://www.w3.org/2000/svg}"
    symbols = {n.get("id"): n for n in root.iter() if n.get("id")}

    def expand(parent, chain=()):
        for node in list(parent):
            if node.tag != ns + "use":
                expand(node, chain)
                continue
            href = node.get("href", node.get("{http://www.w3.org/1999/xlink}href", ""))
            if not href.startswith("#") or href[1:] not in symbols:
                raise ValueError(f"Unresolved SVG use reference: {href}")
            ref = href[1:]
            if ref in chain:
                raise ValueError(f"Cyclic SVG use reference: {ref}")
            group = etree.Element(ns + "g")
            for key, value in node.attrib.items():
                if key not in (
                    "x",
                    "y",
                    "href",
                    "transform",
                    "{http://www.w3.org/1999/xlink}href",
                ):
                    group.set(key, value)
            group.set(
                "transform",
                (
                    node.get("transform", "")
                    + f" translate({node.get('x', '0')},{node.get('y', '0')})"
                ).strip(),
            )
            target = deepcopy(symbols[ref])
            for descendant in target.iter():
                descendant.attrib.pop("id", None)
            group.append(target)
            expand(group, chain + (ref,))
            parent.replace(node, group)

    expand(root)
