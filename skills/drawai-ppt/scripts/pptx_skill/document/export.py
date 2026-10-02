"""Build editable OOXML objects from the portable document model."""

from __future__ import annotations
import base64
from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import tempfile
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData, XyChartData, BubbleChartData
from pptx.util import Pt
from ..io import validate_pptx_package, copy_slide
from .model import load_document
from .xml import (
    EMU,
    FILL,
    SP_ORDER,
    asset_data,
    el,
    ensure,
    ordered,
    replace,
    set_paint,
    stroke,
    tag,
)
from .rich import geometry, hyperlink, shape_effects, text_body, three_d
from .advanced import object_extensions, picture_options, theme_properties
from .charts import chart_details
from .identity import write_identity
from .editing import clear_overrides
from .package import apply_bindings

CHART_TYPES = {
    "barChart": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "lineChart": XL_CHART_TYPE.LINE,
    "areaChart": XL_CHART_TYPE.AREA,
    "pieChart": XL_CHART_TYPE.PIE,
    "doughnutChart": XL_CHART_TYPE.DOUGHNUT,
    "scatterChart": XL_CHART_TYPE.XY_SCATTER,
    "bubbleChart": XL_CHART_TYPE.BUBBLE,
    "radarChart": XL_CHART_TYPE.RADAR,
}


class ExportContext:
    def __init__(self, doc, root, profile):
        self.doc = doc
        self.root = root
        self.profile = profile
        self.issues = []
        self.used = set()
        self.objects = {}

    def audit(self, where, props, used):
        self.used.update(used)
        bound = {b.property_id for b in self.doc.native_bindings if b.owner_id == where}
        for key in props.keys() - used - bound:
            self.issues.append(
                {
                    "object_id": where,
                    "property": key,
                    "requested": props[key],
                    "status": "unsupported",
                    "reason": "No native authoring binding; preserve it through an imported package or implement the binding.",
                }
            )

    def asset(self, ref):
        return asset_data(self.doc.assets[ref], self.root)


def frame(props, *, allow_zero=False):
    v = props.get("object.frame", {"left": 0, "top": 0, "width": 100, "height": 100})
    invalid = (
        (min(v["width"], v["height"]) < 0 or max(v["width"], v["height"]) == 0)
        if allow_zero
        else (v["width"] <= 0 or v["height"] <= 0)
    )
    if invalid:
        raise ValueError("Frame width and height must be positive points")
    return [Pt(v[k]) for k in ("left", "top", "width", "height")]


def common(shape, props, ctx):
    used = set()
    sp = shape._element.find(tag("p:spPr"))
    if sp is None:
        sp = shape._element.find(tag("p:grpSpPr"))
    if "object.frame" in props:
        shape.left, shape.top, shape.width, shape.height = frame(
            props, allow_zero=shape._element.tag == tag("p:cxnSp")
        )
        used.add("object.frame")
    if "object.name" in props:
        shape.name = props["object.name"]
        used.add("object.name")
    if sp is not None:
        xfrm = ensure(sp, "a:xfrm")
        for full, name, mul in [
            ("object.rotation", "rot", 60000),
            ("object.flip_horizontal", "flipH", None),
            ("object.flip_vertical", "flipV", None),
        ]:
            if full in props:
                v = props[full]
                xfrm.set(name, str(round(v * mul)) if mul else ("1" if v else "0"))
                used.add(full)
        if "object.group_child_frame" in props:
            v = props["object.group_child_frame"]
            replace(
                xfrm,
                "a:chOff",
                el(
                    "a:chOff", {"x": round(v["left"] * EMU), "y": round(v["top"] * EMU)}
                ),
            )
            replace(
                xfrm,
                "a:chExt",
                el(
                    "a:chExt",
                    {"cx": round(v["width"] * EMU), "cy": round(v["height"] * EMU)},
                ),
            )
            used.add("object.group_child_frame")
        if "object.black_white_mode" in props:
            sp.set("bwMode", props["object.black_white_mode"])
            used.add("object.black_white_mode")
        used.update(geometry(sp, props))
        used.update(shape_effects(sp, props))
        used.update(three_d(sp, props))
        if "paint.kind" in props:
            value = paint_props(props)
            set_paint(sp, value, shape.part, ctx.doc.assets, ctx.root)
            used.add("paint.kind")
            used.update(
                k for k in props if k.startswith("paint." + props["paint.kind"] + ".")
            )
            if props["paint.kind"] == "solid":
                used.update(k for k in props if k.startswith("color."))
        if any(k.startswith("stroke.") for k in props):
            value = collect(props, "stroke.")
            replace(sp, "a:ln", stroke(value, shape.part, ctx.doc.assets, ctx.root))
            used.update(k for k in props if k.startswith("stroke."))
        ordered(sp, SP_ORDER)
    c = shape._element.find(".//" + tag("p:cNvPr"))
    if c is not None:
        for key, name in [
            ("object.hidden", "hidden"),
            ("accessibility.title", "title"),
            ("accessibility.description", "descr"),
        ]:
            if key in props:
                c.set(
                    name,
                    "1"
                    if props[key] is True
                    else "0"
                    if props[key] is False
                    else str(props[key]),
                )
                used.add(key)
        for key, name in [("interaction.hyperlink", "a:hlinkClick")]:
            if key in props:
                hyperlink(c, name, props[key], shape.part)
                used.add(key)
        if "interaction.navigation_action" in props:
            action = {
                "next_slide": "nextslide",
                "previous_slide": "previousslide",
                "first_slide": "firstslide",
                "last_slide": "lastslide",
                "last_viewed_slide": "lastslideviewed",
                "end_show": "endshow",
            }
            key = props["interaction.navigation_action"]
            if key in action:
                hyperlink(
                    c,
                    "a:hlinkClick",
                    {"action": "ppaction://hlinkshowjump?jump=" + action[key]},
                    shape.part,
                )
                used.add("interaction.navigation_action")
    if getattr(shape, "has_text_frame", False):
        used.update(
            text_body(
                shape.text_frame._txBody, props, shape.part, ctx.doc.assets, ctx.root
            )
        )
    if "object.lock_aspect_ratio" in props or "object.locks" in props:
        nv = shape._element.find(tag("p:nvSpPr")) or shape._element.find(
            tag("p:nvPicPr")
        )
        if nv is not None:
            cnv = ensure(nv, "p:cNvSpPr" if nv.tag == tag("p:nvSpPr") else "p:cNvPicPr")
            locks = ensure(
                cnv, "a:spLocks" if nv.tag == tag("p:nvSpPr") else "a:picLocks"
            )
            if "object.lock_aspect_ratio" in props:
                locks.set(
                    "noChangeAspect", "1" if props["object.lock_aspect_ratio"] else "0"
                )
                used.add("object.lock_aspect_ratio")
            if "object.locks" in props:
                for k, v in props["object.locks"].items():
                    locks.set(k, "1" if v else "0")
                used.add("object.locks")
    if "object.id" in props:
        used.add("object.id")
    if "object.kind" in props:
        used.add("object.kind")
    if "object.z_order" in props:
        used.add("object.z_order")  # sorted before creation
    used.update(object_extensions(shape, props))
    return used


def collect(props, prefix):
    result = {}
    for k, v in props.items():
        if not k.startswith(prefix):
            continue
        keys = k[len(prefix) :].split(".")
        cursor = result
        for key in keys[:-1]:
            cursor = cursor.setdefault(key, {})
        cursor[keys[-1]] = v
    return result


def paint_props(props):
    value = collect(props, "paint.")
    kind = value["kind"]
    if kind == "solid":
        v = {"kind": kind, **value.get("solid", {})}
        channels = collect(props, "color.")
        if channels:
            base = v.get("color", {})
            if isinstance(base, str):
                base = {"kind": "rgb", "channels": base}
            v["color"] = {**base, **channels}
        return v
    if kind == "gradient":
        gradient = value.get("gradient", {})
        return {
            "kind": kind,
            "gradient_kind": gradient.get("kind", gradient.get("path_kind", "linear")),
            **{k: v for k, v in gradient.items() if k not in ("kind", "path_kind")},
        }
    if kind in ("pattern", "image"):
        return {"kind": kind, **value.get(kind, {})}
    return {"kind": kind}


def chart_data(props):
    kind = props.get("chart.type", "barChart")
    series = props["chart.series"]
    data = (
        BubbleChartData()
        if kind == "bubbleChart"
        else XyChartData()
        if kind == "scatterChart"
        else CategoryChartData()
    )
    if kind not in ("scatterChart", "bubbleChart"):
        data.categories = props.get("chart.categories", [])
    for s in series:
        if kind in ("scatterChart", "bubbleChart"):
            target = data.add_series(s["name"])
            for point in s["points"]:
                target.add_data_point(*point)
        else:
            data.add_series(s["name"], s["values"])
    return data


def chart_options(chart, props, ctx):
    used = {"chart.type", "chart.series", "chart.categories"} & props.keys()
    # Actual namespace is schemas.openxmlformats.org, retained from chart root.
    cns = chart._chartSpace.nsmap["c"]

    def ctag(name):
        return "{" + cns + "}" + name

    def ce(name, attrs=None):
        n = etree.Element(ctag(name))
        for k, v in (attrs or {}).items():
            n.set(k, str(v))
        return n

    tree = chart._chartSpace
    chartnode = tree.find(ctag("chart"))
    plot = chartnode.find(ctag("plotArea"))
    if "chart.title" in props:
        value = props["chart.title"]
        value = {"text": value} if isinstance(value, str) else value
        chart.has_title = True
        chart.chart_title.text_frame.text = value.get("text", "")
        style = value.get("properties", {})
        if style:
            text_body(
                chart.chart_title.text_frame._txBody,
                style,
                chart.part,
                ctx.doc.assets,
                ctx.root,
            )
        used.add("chart.title")
    else:
        auto = chartnode.find(ctag("autoTitleDeleted"))
        if auto is not None:
            auto.set("val", "1")
    if "chart.legend" in props:
        v = props["chart.legend"]
        chart.has_legend = v.get("visible", True)
        if chart.has_legend:
            chart.legend.position = {
                "bottom": XL_LEGEND_POSITION.BOTTOM,
                "top": XL_LEGEND_POSITION.TOP,
                "left": XL_LEGEND_POSITION.LEFT,
                "right": XL_LEGEND_POSITION.RIGHT,
            }[v.get("position", "bottom")]
            chart.legend.include_in_layout = v.get("include_in_layout", False)
        used.add("chart.legend")
    for key, name in [
        ("chart.visible_data_only", "plotVisOnly"),
        ("chart.blank_values", "dispBlanksAs"),
    ]:
        if key in props:
            old = chartnode.find(ctag(name))
            if old is not None:
                chartnode.remove(old)
            chartnode.append(
                ce(
                    name,
                    {
                        "val": int(props[key])
                        if isinstance(props[key], bool)
                        else props[key]
                    },
                )
            )
            used.add(key)
    if "chart.style_ref" in props:
        chart.style = int(props["chart.style_ref"])
        used.add("chart.style_ref")
    if "chart.axes" in props:
        axes = (
            plot.findall(ctag("catAx"))
            + plot.findall(ctag("valAx"))
            + plot.findall(ctag("dateAx"))
        )
        for spec in props["chart.axes"]:
            ax = axes[spec.get("index", 0)]
            if "scale" in spec:
                scaling = ax.find(ctag("scaling"))
                for k, v in spec["scale"].items():
                    native = {
                        "minimum": "min",
                        "maximum": "max",
                        "log_base": "logBase",
                        "orientation": "orientation",
                    }[k]
                    old = scaling.find(ctag(native))
                    if old is not None:
                        scaling.remove(old)
                    scaling.append(ce(native, {"val": v}))
            if "number_format" in spec:
                old = ax.find(ctag("numFmt"))
                if old is not None:
                    ax.remove(old)
                ax.append(
                    ce(
                        "numFmt",
                        {"formatCode": spec["number_format"], "sourceLinked": 0},
                    )
                )
            for k, native in [
                ("major_tick", "majorTickMark"),
                ("minor_tick", "minorTickMark"),
                ("label_position", "tickLblPos"),
                ("crosses", "crosses"),
                ("crosses_at", "crossesAt"),
                ("major_unit", "majorUnit"),
                ("minor_unit", "minorUnit"),
            ]:
                if k in spec:
                    old = ax.find(ctag(native))
                    if old is not None:
                        ax.remove(old)
                    ax.append(ce(native, {"val": spec[k]}))
        used.add("chart.axes")
    for key, name, mul in [
        ("chart.series.gap_overlap", "gapWidth", None),
        ("chart.series.doughnut_hole", "holeSize", 100),
    ]:
        if key in props:
            value = props[key]
            if isinstance(value, dict):
                value = value.get("gap_width", 150)
            node = next((c for c in plot if c.tag.endswith("Chart")), None)
            old = node.find(ctag(name))
            if old is not None:
                node.remove(old)
            node.append(ce(name, {"val": round(value * mul) if mul else value}))
            used.add(key)
    series = tree.findall(".//" + ctag("ser"))
    if "chart.series.style" in props:
        for spec in props["chart.series.style"]:
            ser = series[spec["index"]]
            sp = ser.find(ctag("spPr"))
            if sp is None:
                sp = ce("spPr")
                ser.append(sp)
            if "fill" in spec:
                set_paint(sp, spec["fill"], chart.part, ctx.doc.assets, ctx.root)
            if "stroke" in spec:
                replace(
                    sp,
                    "a:ln",
                    stroke(spec["stroke"], chart.part, ctx.doc.assets, ctx.root),
                )
        used.add("chart.series.style")
    if "chart.series.smoothing" in props:
        for ser in series:
            old = ser.find(ctag("smooth"))
            if old is not None:
                ser.remove(old)
            ser.append(ce("smooth", {"val": int(props["chart.series.smoothing"])}))
        used.add("chart.series.smoothing")
    for ser in series:
        ordered(
            ser,
            [
                "c:idx",
                "c:order",
                "c:tx",
                "c:spPr",
                "c:invertIfNegative",
                "c:pictureOptions",
                "c:marker",
                "c:dPt",
                "c:dLbls",
                "c:trendline",
                "c:errBars",
                "c:cat",
                "c:xVal",
                "c:val",
                "c:yVal",
                "c:bubbleSize",
                "c:smooth",
                "c:shape",
                "c:extLst",
            ],
        )
    ordered(
        chartnode,
        [
            "c:title",
            "c:autoTitleDeleted",
            "c:pivotFmts",
            "c:view3D",
            "c:floor",
            "c:sideWall",
            "c:backWall",
            "c:plotArea",
            "c:legend",
            "c:plotVisOnly",
            "c:dispBlanksAs",
            "c:showDLblsOverMax",
            "c:extLst",
        ],
    )
    return used


def table_options(shape, props, ctx):
    table = shape.table
    used = set()
    for key in (
        "first_row",
        "last_row",
        "first_column",
        "last_column",
        "band_rows",
        "band_columns",
    ):
        full = "table." + key
        if full in props:
            setattr(
                table,
                key
                if key not in ("band_rows", "band_columns")
                else {"band_rows": "horz_banding", "band_columns": "vert_banding"}[key],
                props[full],
            )
            used.add(full)
    if "table.grid_columns" in props:
        for column, width in zip(
            table.columns, props["table.grid_columns"], strict=True
        ):
            column.width = Pt(width)
        used.add("table.grid_columns")
    if "table.rows.height" in props:
        for row, height in zip(table.rows, props["table.rows.height"], strict=True):
            row.height = Pt(height)
        used.add("table.rows.height")
    cells = props.get("table.cells.content", [])
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            if isinstance(value, str):
                cell.text = value
            else:
                text_body(
                    cell.text_frame._txBody, value, shape.part, ctx.doc.assets, ctx.root
                )
    if cells:
        used.add("table.cells.content")
    if "table.cells.merge" in props:
        for m in props["table.cells.merge"]:
            table.cell(m["row"], m["column"]).merge(
                table.cell(
                    m["row"] + m["row_span"] - 1, m["column"] + m["column_span"] - 1
                )
            )
        used.add("table.cells.merge")
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cp = cell._tc.get_or_add_tcPr()
            if "table.cells.fill" in props:
                v = props["table.cells.fill"]
                v = v[r][c] if isinstance(v, list) else v
                set_paint(cp, v, shape.part, ctx.doc.assets, ctx.root)
                used.add("table.cells.fill")
            if "table.cells.borders" in props:
                for side, spec in props["table.cells.borders"].items():
                    ln = stroke(spec, shape.part, ctx.doc.assets, ctx.root)
                    ln.tag = tag(
                        "a:"
                        + {
                            "left": "lnL",
                            "right": "lnR",
                            "top": "lnT",
                            "bottom": "lnB",
                            "diagonal_down": "lnTlToBr",
                            "diagonal_up": "lnBlToTr",
                        }[side]
                    )
                    replace(cp, "a:" + etree.QName(ln).localname, ln)
                used.add("table.cells.borders")
            if "table.cells.insets" in props:
                for key, value in props["table.cells.insets"].items():
                    cp.set(
                        {
                            "left": "marL",
                            "right": "marR",
                            "top": "marT",
                            "bottom": "marB",
                        }[key],
                        str(round(value * EMU)),
                    )
                used.add("table.cells.insets")
            for key, name in [
                ("vertical_anchor", "anchor"),
                ("text_direction", "vert"),
                ("anchor_center", "anchorCtr"),
                ("overflow", "horzOverflow"),
            ]:
                full = "table.cells." + key
                if full in props:
                    cp.set(
                        name,
                        "1"
                        if props[full] is True
                        else "0"
                        if props[full] is False
                        else str(props[full]),
                    )
                    used.add(full)
            ordered(
                cp,
                [
                    "a:lnL",
                    "a:lnR",
                    "a:lnT",
                    "a:lnB",
                    "a:lnTlToBr",
                    "a:lnBlToTr",
                    "a:cell3D",
                    *FILL,
                    "a:headers",
                    "a:extLst",
                ],
            )
    if "table.style_ref" in props:
        node = table._tbl.tblPr.find(tag("a:tableStyleId"))
        if node is None:
            node = el("a:tableStyleId")
            table._tbl.tblPr.append(node)
        node.text = props["table.style_ref"]
        used.add("table.style_ref")
    if "table.rtl" in props:
        table._tbl.tblPr.set("rtl", "1" if props["table.rtl"] else "0")
        used.add("table.rtl")
    return used


def add_object(container, obj, ctx, existing=None):
    props = obj.properties
    if existing is not None:
        original_kind = (obj.baseline or {}).get("object.kind")
        if original_kind is not None and original_kind != obj.kind:
            raise ValueError(
                f"{obj.id}: create a replacement object to change kind from {original_kind} to {obj.kind}"
            )
        shape = existing
        if obj.baseline is not None:
            props = {k: v for k, v in props.items() if obj.baseline.get(k) != v}
    elif obj.native_xml:
        node = etree.fromstring(
            obj.native_xml.encode(),
            etree.XMLParser(resolve_entities=False, no_network=True),
        )
        container.shapes._spTree.insert_element_before(node, "p:extLst")
        shape = container.shapes[-1]
    elif obj.kind == "group":
        shape = container.shapes.add_group_shape()
    elif obj.kind == "text_box":
        shape = container.shapes.add_textbox(*frame(props))
    elif obj.kind == "connector":
        x, y, w, h = frame(props, allow_zero=True)
        shape = container.shapes.add_connector(
            {
                "straight": MSO_CONNECTOR.STRAIGHT,
                "elbow": MSO_CONNECTOR.ELBOW,
                "curved": MSO_CONNECTOR.CURVE,
            }[props.get("connector.geometry", "straight")],
            x,
            y,
            x + w,
            y + h,
        )
    elif obj.kind == "picture":
        shape = container.shapes.add_picture(
            BytesIO(ctx.asset(props["picture.asset_ref"])), *frame(props)
        )
    elif obj.kind == "table":
        rows = len(props.get("table.cells.content", props.get("table.rows.height", [])))
        columns = len(props.get("table.grid_columns", [])) or len(
            props["table.cells.content"][0]
        )
        shape = container.shapes.add_table(rows, columns, *frame(props))
    elif obj.kind == "chart":
        shape = container.shapes.add_chart(
            CHART_TYPES[props.get("chart.type", "barChart")],
            *frame(props),
            chart_data(props),
        )
    elif obj.kind == "shape":
        shape = container.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, *frame(props))
    elif obj.kind == "ole":
        icon = (
            BytesIO(ctx.asset(props["embedded.preview_ref"]))
            if "embedded.preview_ref" in props
            else None
        )
        shape = container.shapes.add_ole_object(
            BytesIO(ctx.asset(props["embedded.asset_ref"])),
            props["embedded.prog_id"],
            *frame(props),
            icon_file=icon,
        )
    elif obj.kind == "video":
        poster = (
            BytesIO(ctx.asset(props["media.poster_ref"]))
            if "media.poster_ref" in props
            else None
        )
        shape = container.shapes.add_movie(
            BytesIO(ctx.asset(props["media.asset_ref"])),
            *frame(props),
            poster_frame_image=poster,
            mime_type=ctx.doc.assets[props["media.asset_ref"]].get(
                "content_type", "video/mp4"
            ),
        )
    else:
        raise ValueError(
            f"{obj.id}: {obj.kind} needs a native_xml object and package relationships; cannot create it faithfully from incomplete settings"
        )
    if existing is None and not obj.native_xml and "object.style_ref" not in props:
        style = shape._element.find(tag("p:style"))
        if style is not None:
            effect = style.find(tag("a:effectRef"))
            if effect is not None:
                effect.set("idx", "0")
    ctx.objects[obj.id] = shape
    identity = shape._element.find(".//" + tag("p:cNvPr"))
    if identity is not None:
        write_identity(identity, obj.id)
    used = common(shape, props, ctx)
    if obj.kind == "group":
        original_transform = (
            deepcopy(shape._element.grpSpPr.xfrm)
            if existing is not None and shape._element.grpSpPr.xfrm is not None
            else None
        )
        source = {s.shape_id: s for s in shape.shapes}
        keep = {c.source_shape_id for c in obj.children if c.source_shape_id}
        if existing is not None:
            for child in list(shape.shapes):
                if child.shape_id not in keep:
                    child._element.getparent().remove(child._element)
        children = sorted(
            obj.children, key=lambda child: child.properties.get("object.z_order", 0)
        )
        child_shapes = [
            add_object(shape, child, ctx, source.get(child.source_shape_id))
            for child in children
        ]
        reorder_shapes(shape.shapes._spTree, [child._element for child in child_shapes])
        if original_transform is not None:
            replace(shape._element.grpSpPr, "a:xfrm", original_transform)
        # add_group_shape recalculates bounds while children are added.
        used.update(
            common(
                shape, {k: v for k, v in props.items() if k.startswith("object.")}, ctx
            )
        )
    if obj.kind == "table" and any(k.startswith("table.") for k in props):
        used.update(table_options(shape, props, ctx))
    if obj.kind == "chart" and (
        existing is None or any(k.startswith("chart.") for k in props)
    ):
        if existing is not None and any(
            k in props for k in ("chart.series", "chart.categories")
        ):
            shape.chart.replace_data(chart_data(obj.properties))
        if (
            existing is not None
            and "chart.type" in props
            and obj.baseline
            and props["chart.type"] != obj.baseline.get("chart.type")
        ):
            raise ValueError(
                "Changing an imported chart family requires creating a replacement chart or editing its exact plot XML"
            )
        used.update(chart_options(shape.chart, props, ctx))
        used.update(chart_details(shape.chart, props, ctx))
    if obj.kind == "connector":
        used.update({"connector.geometry"} & props.keys())
    if obj.kind == "picture":
        used.update({"picture.asset_ref"} & props.keys())
        if existing is not None and "picture.asset_ref" in props:
            _, rid = shape.part.get_or_add_image_part(
                BytesIO(ctx.asset(props["picture.asset_ref"]))
            )
            shape._element.blipFill.blip.rEmbed = rid
        used.update(picture_options(shape, props, ctx))
        if "picture.crop" in props:
            values = {}
            for key, value in props["picture.crop"].items():
                side = {
                    "l": "left",
                    "r": "right",
                    "t": "top",
                    "b": "bottom",
                    "left": "left",
                    "right": "right",
                    "top": "top",
                    "bottom": "bottom",
                }.get(key)
                if side is None:
                    raise ValueError(f"Unknown picture crop edge: {key}")
                if side in values and values[side] != value:
                    raise ValueError(f"Conflicting picture crop values for {side}")
                values[side] = value
            for side, value in values.items():
                setattr(shape, "crop_" + side, value)
            used.add("picture.crop")
    if obj.kind == "ole":
        used.update(
            {"embedded.asset_ref", "embedded.prog_id", "embedded.preview_ref"}
            & props.keys()
        )
    if obj.kind == "video":
        used.update({"media.asset_ref", "media.poster_ref"} & props.keys())
    if obj.clear_overrides:
        clear_overrides(shape, obj.clear_overrides)
    ctx.audit(obj.id, props, used)
    return shape


def reorder_shapes(tree, elements):
    """Reorder modeled drawing nodes in their slots, keeping unknown XML in place."""
    known = set(elements)
    ordered_nodes = iter(elements)
    tree[:] = [next(ordered_nodes) if child in known else child for child in tree]


def slide_properties(slide, props, ctx):
    used = set()
    for key, name in [
        ("slide.hidden", "show"),
        ("slide.show_master_shapes", "showMasterSp"),
    ]:
        if key in props:
            slide._element.set(
                name, str(int(not props[key] if key == "slide.hidden" else props[key]))
            )
            used.add(key)
    if "slide.name" in props:
        slide._element.cSld.set("name", props["slide.name"])
        used.add("slide.name")
    if "slide.background" in props:
        bg = slide._element.cSld.get_or_add_bgPr()
        set_paint(bg, props["slide.background"], slide.part, ctx.doc.assets, ctx.root)
        ordered(bg, [*FILL, "a:effectLst", "a:extLst"])
        used.add("slide.background")
    if "notes.content" in props:
        v = props["notes.content"]
        slide.notes_slide.notes_text_frame.text = v if isinstance(v, str) else v["text"]
        used.add("notes.content")
    transition = {
        k.split(".", 1)[1]: v for k, v in props.items() if k.startswith("transition.")
    }
    if transition:
        node = el(
            "p:transition",
            {
                "spd": transition.get("speed"),
                "advClick": transition.get("advance_on_click", True),
                "advTm": transition.get("advance_after"),
            },
        )
        kind = transition.get("kind", "fade")
        if kind not in (
            "cut",
            "fade",
            "push",
            "wipe",
            "split",
            "cover",
            "pull",
            "randomBar",
            "wheel",
            "strips",
            "zoom",
            "circle",
            "diamond",
            "plus",
            "wedge",
            "blinds",
            "checker",
            "comb",
            "random",
            "newsflash",
            "dissolve",
        ):
            raise ValueError(
                f"Unsupported native transition {kind}; use preserved native XML for Office extensions"
            )
        a = {}
        if "direction" in transition:
            a["dir"] = transition["direction"]
        node.append(el("p:" + kind, a))
        replace(slide._element, "p:transition", node)
        used.update(
            "transition." + k
            for k in transition
            if k in ("kind", "speed", "advance_on_click", "advance_after", "direction")
        )
    if "animation.timeline" in props:
        v = props["animation.timeline"]
        tree = (
            etree.fromstring(
                v.encode(), etree.XMLParser(resolve_entities=False, no_network=True)
            )
            if isinstance(v, str)
            else native_tree(v)
        )
        if tree.tag != tag("p:timing"):
            raise ValueError("animation.timeline must be a native p:timing tree")
        replace(slide._element, "p:timing", tree)
        used.add("animation.timeline")
    ordered(
        slide._element,
        ["p:cSld", "p:clrMapOvr", "p:transition", "p:timing", "p:extLst"],
    )
    return used


def native_tree(value):
    node = el(value["tag"], value.get("attributes", {}))
    if "text" in value:
        node.text = value["text"]
    node.extend(native_tree(c) for c in value.get("children", []))
    return node


def patch_template_package(original, saved):
    if original is None:
        return saved
    with zipfile.ZipFile(BytesIO(original)) as z:
        before = {n: z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(BytesIO(saved)) as z:
        after = {n: z.read(n) for n in z.namelist()}
    if any(n.startswith("_xmlsignatures/") for n in before):
        raise ValueError(
            "Signed template requires an explicit signature removal workflow before editing"
        )
    # Preserve orphan/custom parts which python-pptx does not walk. Merge their content types.
    out = {**before, **after}
    tree = etree.fromstring(after["[Content_Types].xml"])
    known = {(n.tag, n.get("PartName", n.get("Extension"))) for n in tree}
    for node in etree.fromstring(before["[Content_Types].xml"]):
        key = (node.tag, node.get("PartName", node.get("Extension")))
        if key not in known:
            tree.append(node)
    out["[Content_Types].xml"] = etree.tostring(
        tree, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in out.items():
            z.writestr(n, b)
    return stream.getvalue()


def export_pptx(document, output, *, asset_root=None, profile="strict"):
    if profile not in ("strict", "best_effort"):
        raise ValueError("Unknown export profile")
    doc = load_document(document)
    output = Path(output)
    if output.suffix.lower() != ".pptx":
        raise ValueError(
            "This exporter writes the PPTX container; macro-enabled/encrypted containers need a separate host profile"
        )
    if output.exists():
        raise FileExistsError(output)
    root = Path(
        asset_root
        or (
            Path(document).resolve().parent
            if isinstance(document, (str, Path))
            else "."
        )
    )
    ctx = ExportContext(doc, root, profile)
    original = None
    if doc.template:
        original = (
            base64.b64decode(doc.template["data"], validate=True)
            if "data" in doc.template
            else (root / doc.template["path"]).read_bytes()
        )
        if hashlib.sha256(original).hexdigest() != doc.template["sha256"]:
            raise ValueError("Template SHA256 differs from imported document")
        with zipfile.ZipFile(BytesIO(original)) as z:
            if any(n.endswith("vbaProject.bin") for n in z.namelist()):
                raise ValueError(
                    "Macro-enabled template requires a PPTM profile; do not rename it to PPTX"
                )
    prs = Presentation(BytesIO(original)) if original else Presentation()
    props = doc.properties
    if doc.baseline is not None:
        props = {k: v for k, v in props.items() if doc.baseline.get(k) != v}
    used = theme_properties(prs, props)
    if "document.page_size" in props:
        v = props["document.page_size"]
        prs.slide_width = Pt(v["width"])
        prs.slide_height = Pt(v["height"])
        used.add("document.page_size")
    elif not original:
        prs.slide_width = Pt(960)
        prs.slide_height = Pt(540)
    for key in (
        "title",
        "subject",
        "author",
        "description",
        "revision",
        "keywords",
        "created",
        "modified",
    ):
        full = "document.metadata." + key
        if full in props:
            value = props[full]
            if key == "keywords" and isinstance(value, list):
                value = ", ".join(value)
            if key in ("created", "modified"):
                from datetime import datetime

                value = datetime.fromisoformat(value.replace("Z", "+00:00"))
            if key == "revision":
                node = prs.core_properties._element.find(
                    "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}revision"
                )
                if node is None:
                    node = etree.SubElement(
                        prs.core_properties._element,
                        "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}revision",
                    )
                node.text = str(value)
            else:
                setattr(
                    prs.core_properties,
                    "comments" if key == "description" else key,
                    value,
                )
            used.add(full)
    existing = list(prs.slides)
    wanted = []
    for page in doc.pages:
        source = existing[page.source_slide - 1] if page.source_slide else None
        if source is not None and source not in wanted:
            slide = source
        elif source is not None:
            slide = copy_slide(prs, source)
        else:
            layoutref = page.properties.get("slide.layout_ref")
            layouts = [l for m in prs.slide_masters for l in m.slide_layouts]
            layout = (
                next(
                    (
                        l
                        for l in layouts
                        if str(l.part.partname) == layoutref or l.name == layoutref
                    ),
                    None,
                )
                if layoutref
                else prs.slide_layouts[6]
            )
            if layout is None:
                raise ValueError(f"Layout not found: {layoutref}")
            slide = prs.slides.add_slide(layout)
        wanted.append(slide)
        write_identity(slide._element.cSld, page.id, "p:extLst")
        pp = page.properties
        if page.baseline is not None:
            pp = {k: v for k, v in pp.items() if page.baseline.get(k) != v}
        if source is not None and "slide.layout_ref" in pp:
            raise ValueError(
                "Changing an imported page layout needs explicit placeholder remapping; use the plan add_slide/duplicate_slide workflow or create a new bound page"
            )
        sused = slide_properties(slide, pp, ctx)
        if "slide.layout_ref" in pp:
            sused.add("slide.layout_ref")
        ctx.audit(page.id, pp, sused)
        shape_map = {s.shape_id: s for s in slide.shapes}
        keep = {o.source_shape_id for o in page.objects if o.source_shape_id}
        if source:
            for s in list(slide.shapes):
                if s.shape_id not in keep:
                    s._element.getparent().remove(s._element)
        ordered_objects = sorted(
            page.objects, key=lambda o: o.properties.get("object.z_order", 0)
        )
        rendered_shapes = [
            add_object(slide, obj, ctx, shape_map.get(obj.source_shape_id))
            for obj in ordered_objects
        ]
        reorder_shapes(
            slide.shapes._spTree, [shape._element for shape in rendered_shapes]
        )
        # Resolve connections after every object has its real numeric shape ID.
        for obj in page.objects:
            if obj.kind != "connector":
                continue
            shape = ctx.objects[obj.id]
            for key, method in [
                ("connector.start_target", "begin_connect"),
                ("connector.end_target", "end_connect"),
            ]:
                if key in obj.properties:
                    value = obj.properties[key]
                    getattr(shape, method)(
                        ctx.objects[value["object_ref"]], value.get("site", 0)
                    )
                    ctx.issues = [
                        i
                        for i in ctx.issues
                        if not (i["object_id"] == obj.id and i["property"] == key)
                    ]
                    ctx.used.add(key)
    # Order/remove pages while retaining the actual source page relationships.
    ids = prs.slides._sldIdLst
    by_slide = {prs.part.related_slide(n.rId).part: n for n in ids}
    ids[:] = [by_slide[s.part] for s in wanted]
    for part, node in by_slide.items():
        if part not in {s.part for s in wanted}:
            prs.part.drop_rel(node.rId)
    identifier = prs.core_properties._element.find(
        "{http://purl.org/dc/elements/1.1/}identifier"
    )
    if identifier is None:
        identifier = etree.SubElement(
            prs.core_properties._element, "{http://purl.org/dc/elements/1.1/}identifier"
        )
    identifier.text = doc.id
    ctx.audit(doc.id, props, used)
    if ctx.issues and profile == "strict":
        raise ValueError(
            "Unsupported properties: " + json.dumps(ctx.issues, ensure_ascii=False)
        )
    stream = BytesIO()
    prs.save(stream)
    data = patch_template_package(original, stream.getvalue())
    if doc.native_parts:
        with zipfile.ZipFile(BytesIO(data)) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        types = etree.fromstring(parts["[Content_Types].xml"])
        ns = "{http://schemas.openxmlformats.org/package/2006/content-types}"
        for path, spec in doc.native_parts.items():
            if path.endswith("vbaProject.bin"):
                raise ValueError(
                    "Macro project cannot be inserted into the PPTX profile"
                )
            if path.startswith("/") or ".." in Path(path).parts:
                raise ValueError(f"Invalid package part path {path}")
            parts[path] = base64.b64decode(spec["data"], validate=True)
            if not any(n.get("PartName") == "/" + path for n in types):
                types.append(
                    etree.Element(
                        ns + "Override",
                        PartName="/" + path,
                        ContentType=spec["content_type"],
                    )
                )
        parts["[Content_Types].xml"] = etree.tostring(
            types, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        stream = BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as z:
            for n, b in parts.items():
                z.writestr(n, b)
        data = stream.getvalue()
    data, binding_report = apply_bindings(data, doc.native_bindings)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        temp = Path(tmp) / "check.pptx"
        temp.write_bytes(data)
        validation = validate_pptx_package(temp)
        Presentation(str(temp))
        with output.open("xb") as f:
            f.write(data)
    return {
        "schema": "presentation.export_report.v1",
        "format": "pptx",
        "output": str(output.resolve()),
        "profile": profile,
        "properties_written": sorted(ctx.used),
        "native_bindings": binding_report,
        "issues": ctx.issues,
        "validation": validation,
        "native_objects": len(ctx.objects),
    }
