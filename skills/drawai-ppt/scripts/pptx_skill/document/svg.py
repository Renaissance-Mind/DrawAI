"""Static SVG projection using the same office layout as the native PPTX."""

from __future__ import annotations
import json
import base64
import zlib
import hashlib
from pathlib import Path
import shutil
import subprocess
import tempfile
from lxml import etree
from pptx import Presentation
from .model import load_document
from .export import export_pptx
from .render_env import office_environment
from .xml import asset_data
from .svg_fallback import path_gradient_objects, apply_gradient_fallback

SVG = "{http://www.w3.org/2000/svg}"


def export_svg(
    document,
    output_dir,
    *,
    asset_root=None,
    renderer=None,
    rasterizer=None,
    profile="strict_native",
):
    doc = load_document(document)
    output = Path(output_dir).resolve()
    root = Path(
        asset_root
        or (
            Path(document).resolve().parent
            if isinstance(document, (str, Path))
            else "."
        )
    )
    if output.exists() and any(output.iterdir()):
        raise FileExistsError("SVG output directory must be empty")
    soffice = renderer or shutil.which("soffice")
    if not soffice:
        raise ValueError(
            "SVG export needs LibreOffice soffice; install LibreOffice or pass --renderer"
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    pages = []
    issues = []
    with tempfile.TemporaryDirectory(prefix="ppt-svg-", dir=output.parent) as tmpname:
        tmp = Path(tmpname)
        ppt = tmp / "document.pptx"
        report = export_pptx(doc, ppt, asset_root=root, profile=profile)
        issues.extend(report["issues"])
        prs = Presentation(str(ppt))
        environment, fonts = office_environment(prs, tmp)
        previews = None
        if any(path_gradient_objects(page) for page in doc.pages):
            from ..engine import render_pptx

            if rasterizer or shutil.which("pdftoppm"):
                previews = render_pptx(
                    ppt,
                    tmp / "gradient-preview",
                    renderer=soffice,
                    rasterizer=rasterizer,
                    dpi=192,
                )["slides"]
            else:
                issues.append(
                    {
                        "property": "paint.gradient.kind",
                        "status": "approximation",
                        "reason": "Poppler is unavailable; rect/shape gradients retain the office SVG approximation.",
                    }
                )
        for i, page in enumerate(doc.pages):
            single = Presentation(str(ppt))
            lst = single.slides._sldIdLst
            for j, node in reversed(list(enumerate(lst))):
                if j != i:
                    single.part.drop_rel(node.rId)
                    lst.remove(node)
            path = tmp / f"slide-{i + 1:03}.pptx"
            single.save(path)
            renderdir = tmp / f"render-{i}"
            renderdir.mkdir()
            subprocess.run(
                [
                    soffice,
                    f"-env:UserInstallation={(tmp / 'profile').as_uri()}",
                    "--headless",
                    "--convert-to",
                    "svg:impress_svg_Export",
                    "--outdir",
                    str(renderdir),
                    str(path),
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=180,
                env=environment,
            )
            svg = renderdir / (path.stem + ".svg")
            if not svg.exists():
                raise ValueError(
                    "LibreOffice produced no SVG; its Impress SVG export filter is required"
                )
            tree = etree.fromstring(
                svg.read_bytes(),
                etree.XMLParser(resolve_entities=False, no_network=True),
            )
            for dummy in tree.xpath(
                '//*[contains(concat(" ",normalize-space(@class)," ")," DummySlide ")]'
            ):
                dummy.getparent().remove(dummy)
            actual = tree.xpath(
                '//*[local-name()="g" and @class="Slide" and @id!="dummy-slide"]'
            )
            if len(actual) != 1:
                raise ValueError(f"Expected one actual SVG slide, found {len(actual)}")
            for node in [actual[0], *actual[0].iterancestors()]:
                if node.get("visibility") == "hidden":
                    node.set("visibility", "visible")
                if node.get("display") == "none":
                    node.attrib.pop("display")
            ooo = "{http://xml.openoffice.org/svg/export}"
            metadata_nodes = [
                n
                for n in tree.iter()
                if n.get(ooo + "slide") == actual[0].get("id") and n.get(ooo + "master")
            ]
            if len(metadata_nodes) != 1:
                raise ValueError(
                    f"SVG export has {len(metadata_nodes)} slide/master metadata records for {actual[0].get('id')}"
                )
            meta = metadata_nodes[0]
            master = meta.get(ooo + "master")
            bg = "bg-" + actual[0].get("id")
            if not tree.xpath("//*[@id=$value]", value=bg):
                bg = "bg-" + master
            sources = []
            if meta.get(ooo + "background-visibility") != "hidden":
                sources.append(bg)
            if meta.get(ooo + "master-objects-visibility") != "hidden":
                sources.append("bo-" + master)
            for index, ref in enumerate(sources):
                if tree.xpath("//*[@id=$value]", value=ref):
                    use = etree.Element(SVG + "use")
                    use.set("{http://www.w3.org/1999/xlink}href", "#" + ref)
                    actual[0].insert(index, use)
            # Deliver a static, self-contained drawing. Drop generated slideshow scripting.
            for n in list(tree.iter()):
                if etree.QName(n).localname in (
                    "script",
                    "animate",
                    "animateTransform",
                    "animateMotion",
                    "set",
                ):
                    n.getparent().remove(n)
                else:
                    for key in list(n.attrib):
                        if key.lower().startswith("on"):
                            del n.attrib[key]
            for image in tree.iter(SVG + "image"):
                href = image.get(
                    "{http://www.w3.org/1999/xlink}href", image.get("href", "")
                )
                if href and not href.startswith(("data:", "#")):
                    raise ValueError(
                        f"SVG exporter emitted a non-embedded resource: {href}"
                    )
            width = prs.slide_width / 12700
            height = prs.slide_height / 12700
            if previews is not None and path_gradient_objects(page):
                issues.extend(
                    apply_gradient_fallback(tree, page, previews[i], width, height)
                )
            tree.set("width", f"{width:g}pt")
            tree.set("height", f"{height:g}pt")
            # Normalize namespace declarations after the filter's nodes were removed.
            tree = etree.fromstring(
                etree.tostring(tree),
                etree.XMLParser(resolve_entities=False, no_network=True),
            )
            portable = doc.model_dump(mode="json", by_alias=True)
            portable["pages"] = [portable["pages"][i]]
            for ref, asset in portable["assets"].items():
                asset["data"] = base64.b64encode(asset_data(asset, root)).decode()
                asset.pop("path", None)
            if portable["template"] and "path" in portable["template"]:
                portable["template"]["data"] = base64.b64encode(
                    (root / portable["template"].pop("path")).read_bytes()
                ).decode()
            packed = base64.b64encode(
                zlib.compress(json.dumps(portable, ensure_ascii=False).encode())
            ).decode()
            digest = hashlib.sha256(etree.tostring(tree, method="c14n")).hexdigest()
            metadata = etree.SubElement(tree, SVG + "metadata")
            metadata.text = json.dumps(
                {
                    "schema": doc.schema_id,
                    "document_id": doc.id,
                    "page_id": page.id,
                    "units": "pt",
                    "projection": "static",
                    "renderer": "LibreOffice",
                    "objects": [o.id for o in page.objects],
                    "visual_sha256": digest,
                    "document_encoding": "zlib-base64",
                    "document_data": packed,
                },
                ensure_ascii=False,
            )
            rendered = etree.tostring(tree, xml_declaration=True, encoding="UTF-8")
            output.mkdir(exist_ok=True)
            target = output / (path.stem + ".svg")
            target.write_bytes(rendered)
            counts = {
                k: sum(1 for _ in tree.iter(SVG + k))
                for k in ("text", "tspan", "path", "image")
            }
            pages.append({"page_id": page.id, "svg": str(target), "elements": counts})
    return {
        "schema": "presentation.export_report.v1",
        "format": "svg",
        "profile": "static_office_projection",
        "renderer": "LibreOffice",
        "fonts": fonts,
        "pages": pages,
        "issues": issues,
        "limits": [
            "SVG is a static visual projection; native chart/table data, animations and media playback remain in the source model/PPTX.",
            "PowerPoint and LibreOffice can differ for Office Math, 3-D, artistic effects and font fallback. SVG text remains text where the export filter supports it.",
        ],
    }
