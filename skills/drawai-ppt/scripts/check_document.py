#!/usr/bin/env python3
"""Real file regression checks for the unified model, runnable outside any repo."""

from __future__ import annotations
import base64
from pathlib import Path
import tempfile
import unittest
import zipfile
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lxml import etree
from pptx import Presentation
from pptx_skill.document.model import Document, PROPERTIES
from pptx_skill.document.export import export_pptx
from pptx_skill.document.importer import import_pptx, import_svg
from pptx_skill.document.xml import tag
from pptx_skill.document.capabilities import capability_report, WRITABLE


def shape(id="s", **properties):
    props = {
        "object.frame": {"left": 20, "top": 30, "width": 180, "height": 80},
        "shape.geometry.preset": "rect",
        "paint.kind": "solid",
        "paint.solid.color": "#DDEEFF",
    }
    props.update(properties)
    if props["paint.kind"] != "solid":
        props.pop("paint.solid.color", None)
    return {"id": id, "properties": props}


def document(objects=None, properties=None):
    return {
        "schema": "presentation.document.v1",
        "properties": {
            "document.page_size": {"width": 960, "height": 540},
            **(properties or {}),
        },
        "pages": [{"id": "page", "objects": objects or [shape()]}],
    }


class Checks(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_rich_text_native_sequence_and_styles(self):
        obj = {
            "id": "text",
            "kind": "text_box",
            "properties": {
                "object.frame": {"left": 30, "top": 30, "width": 500, "height": 200},
                "text.run.font.latin": "Arial",
                "text.run.font.east_asian": "PingFang SC",
                "text.body.autofit": "shrink",
                "text.body.autofit_scale": {"font": 0.9},
                "text.body.insets": {"left": 12},
                "text.body.paragraphs": [
                    {
                        "properties": {
                            "text.paragraph.first_line_indent": -8,
                            "text.paragraph.left_margin": 16,
                            "text.paragraph.space_after": {
                                "mode": "points",
                                "value": 10,
                            },
                            "text.paragraph.line_spacing": {
                                "mode": "ratio",
                                "value": 1.2,
                            },
                            "text.paragraph.tab_stops": [{"position": 40}],
                        },
                        "content": [
                            {
                                "text": "中文",
                                "properties": {
                                    "text.run.font.bold": True,
                                    "text.run.letter_spacing": 0.5,
                                    "text.run.underline.type": "dbl",
                                },
                            },
                            {"kind": "break"},
                            {"kind": "field", "type": "slidenum", "text": "1"},
                            {
                                "text": " tail",
                                "properties": {"text.run.font.italic": True},
                            },
                        ],
                    }
                ],
            },
        }
        path = self.root / "rich.pptx"
        export_pptx(document([obj]), path)
        s = Presentation(path).slides[0].shapes[0]
        self.assertEqual(
            [
                etree.QName(n).localname
                for n in s.text_frame._txBody.findall(tag("a:p"))[0]
            ],
            ["pPr", "r", "br", "fld", "r"],
        )
        self.assertEqual(
            s.text_frame.paragraphs[0]._p.pPr.get("indent"), str(-8 * 12700)
        )
        r = s.text_frame.paragraphs[0].runs[0]
        self.assertTrue(r.font.bold)
        self.assertEqual(r._r.rPr.get("u"), "dbl")

    def test_all_native_shape_and_pattern_enumerations(self):
        objects = [
            shape("preset-" + p, **{"shape.geometry.preset": p})
            for p in PROPERTIES["shape.geometry.preset"]["allowed_values"]
        ]
        objects += [
            shape(
                "pattern-" + p,
                **{
                    "paint.kind": "pattern",
                    "paint.pattern.preset": p,
                    "paint.pattern.foreground": "#336699",
                    "paint.pattern.background": "#FFFFFF",
                },
            )
            for p in PROPERTIES["paint.pattern.preset"]["allowed_values"]
        ]
        path = self.root / "enum.pptx"
        export_pptx(document(objects), path)
        p = Presentation(path)
        self.assertEqual(len(p.slides[0].shapes), 241)
        got = {n.get("prst") for n in p.slides[0]._element.iter(tag("a:prstGeom"))}
        self.assertEqual(
            got, set(PROPERTIES["shape.geometry.preset"]["allowed_values"])
        )

    def test_arrow_dimensions_and_gradient_variants(self):
        objects = []
        for k in PROPERTIES["stroke.end_arrow.kind"]["allowed_values"]:
            for w in ("sm", "med", "lg"):
                for length in ("sm", "med", "lg"):
                    objects.append(
                        shape(
                            f"{k}-{w}-{length}",
                            **{
                                "stroke.start_arrow.kind": k,
                                "stroke.start_arrow.width": w,
                                "stroke.start_arrow.length": length,
                                "stroke.end_arrow.kind": k,
                                "stroke.end_arrow.width": w,
                                "stroke.end_arrow.length": length,
                                "stroke.paint": "#336699",
                            },
                        )
                    )
        for kind in ("linear", "circle", "rect", "shape"):
            objects.append(
                shape(
                    kind,
                    **{
                        "paint.kind": "gradient",
                        "paint.gradient.kind": kind,
                        "paint.gradient.stops": [
                            {"position": 0, "color": "#336699"},
                            {
                                "position": 1,
                                "color": {
                                    "kind": "rgb",
                                    "channels": "FFFFFF",
                                    "alpha": 0.5,
                                },
                            },
                        ],
                    },
                )
            )
        p = self.root / "arrows.pptx"
        export_pptx(document(objects), p)
        s = Presentation(p).slides[0]
        self.assertEqual(len(list(s._element.iter(tag("a:tailEnd")))), 54)
        self.assertEqual(
            {n.get("path") for n in s._element.iter(tag("a:path"))},
            {"circle", "rect", "shape"},
        )

    def test_table_merges_borders_and_native_chart_workbook(self):
        objects = [
            {
                "id": "table",
                "kind": "table",
                "properties": {
                    "object.frame": {
                        "left": 30,
                        "top": 30,
                        "width": 400,
                        "height": 150,
                    },
                    "table.grid_columns": [200, 200],
                    "table.rows.height": [50, 50, 50],
                    "table.cells.content": [["heading", ""], ["a", "1"], ["b", "2"]],
                    "table.cells.merge": [
                        {"row": 0, "column": 0, "row_span": 1, "column_span": 2}
                    ],
                    "table.cells.fill": "#EAF2FF",
                    "table.cells.borders": {"bottom": {"width": 1, "paint": "#336699"}},
                },
            },
            {
                "id": "chart",
                "kind": "chart",
                "properties": {
                    "object.frame": {
                        "left": 30,
                        "top": 200,
                        "width": 600,
                        "height": 250,
                    },
                    "chart.type": "lineChart",
                    "chart.categories": ["a", "b"],
                    "chart.series": [{"name": "values", "values": [1, 2]}],
                    "chart.axes": [{"index": 1, "scale": {"minimum": 0, "maximum": 3}}],
                },
            },
        ]
        path = self.root / "data.pptx"
        export_pptx(document(objects), path)
        prs = Presentation(path)
        self.assertTrue(prs.slides[0].shapes[0].table.cell(0, 0).is_merge_origin)
        self.assertEqual(
            list(prs.slides[0].shapes[1].chart.series[0].values), [1.0, 2.0]
        )
        with zipfile.ZipFile(path) as z:
            self.assertTrue(any(n.endswith(".xlsx") for n in z.namelist()))

    def test_import_edit_preserves_unknown_parts_and_run_style(self):
        p = self.root / "source.pptx"
        obj = {
            "id": "t",
            "kind": "text_box",
            "properties": {
                "object.frame": {"left": 20, "top": 20, "width": 600, "height": 100},
                "text.body.paragraphs": [
                    {
                        "content": [
                            {
                                "text": "first",
                                "properties": {"text.run.font.bold": True},
                            },
                            {
                                "text": " second",
                                "properties": {"text.run.font.italic": True},
                            },
                        ]
                    }
                ],
            },
        }
        source = document([obj])
        source["native_parts"] = {
            "customXml/item99.xml": {
                "content_type": "application/xml",
                "data": base64.b64encode(
                    b'<retained xmlns="urn:custom">actual package extension</retained>'
                ).decode(),
            }
        }
        export_pptx(source, p)
        before = p.read_bytes()
        doc = import_pptx(p)
        doc.pages[0].objects[0].properties["text.body.paragraphs"][0]["content"][0][
            "text"
        ] = "edited"
        doc.save(self.root / "portable.json")
        out = self.root / "edited.pptx"
        export_pptx(self.root / "portable.json", out)
        self.assertEqual(p.read_bytes(), before)
        with zipfile.ZipFile(out) as z:
            self.assertIn(b"actual package extension", z.read("customXml/item99.xml"))
        runs = Presentation(out).slides[0].shapes[0].text_frame.paragraphs[0].runs
        self.assertTrue(runs[0].font.bold)
        self.assertTrue(runs[1].font.italic)
        self.assertEqual(runs[0].text, "edited")

    def test_unmapped_settings_cannot_silently_disappear(self):
        doc = document([shape(**{"smartart.direction": "norm"})])
        out = self.root / "bad.pptx"
        with self.assertRaisesRegex(ValueError, "Unsupported properties"):
            export_pptx(doc, out)
        self.assertFalse(out.exists())
        report = export_pptx(doc, self.root / "compatible.pptx", profile="compatible")
        self.assertEqual(report["issues"][0]["property"], "smartart.direction")
        with self.assertRaisesRegex(ValueError, "Unknown property"):
            Document.model_validate(document([shape(**{"invented.width": 3})]))

    def test_nested_enum_and_shear_rejected(self):
        bad = {
            "id": "t",
            "kind": "text_box",
            "properties": {
                "object.frame": {"left": 1, "top": 1, "width": 10, "height": 10},
                "text.body.paragraphs": [
                    {
                        "content": [
                            {
                                "text": "x",
                                "properties": {"text.run.underline.type": "invented"},
                            }
                        ]
                    }
                ],
            },
        }
        with self.assertRaisesRegex(ValueError, "must be one of"):
            Document.model_validate(document([bad]))
        with self.assertRaisesRegex(ValueError, "shear"):
            export_pptx(
                document([shape(**{"object.transform": [1, 0, 0.5, 1, 0, 0]})]),
                self.root / "shear.pptx",
            )

    def test_prior_vector_formula_tiny_canvas_and_use_paths(self):
        source = Path(__file__).resolve().parents[1] / "assets/demo/formula.svg"
        doc, report = import_svg(source)
        path = self.root / "formula.pptx"
        export_pptx(doc, path)
        prs = Presentation(path)
        self.assertGreaterEqual(prs.slide_height, 914400)
        self.assertGreater(
            len(list(prs.slides[0]._element.iter(tag("a:custGeom")))), 40
        )

    def test_coverage_has_exact_catalog_rows_and_separate_status(self):
        report = capability_report()
        self.assertEqual(report["property_count"], 475)
        self.assertEqual({p["id"] for p in report["properties"]}, set(PROPERTIES))
        self.assertFalse(WRITABLE - set(PROPERTIES))
        self.assertEqual(report["counts"]["native_authoring"], len(WRITABLE))

    def test_horizontal_connector_keeps_exact_zero_height(self):
        obj = {
            "id": "line",
            "kind": "connector",
            "properties": {
                "object.frame": {"left": 10, "top": 20, "width": 120, "height": 0},
                "stroke.paint": "#246",
                "stroke.end_arrow.kind": "triangle",
            },
        }
        path = self.root / "line.pptx"
        export_pptx(document([obj]), path)
        shape = Presentation(path).slides[0].shapes[0]
        self.assertEqual(shape.height, 0)

    @unittest.skipUnless(
        __import__("shutil").which("soffice")
        and __import__("shutil").which("pdftoppm"),
        "Gradient fallback requires actual office and PDF rendering",
    )
    def test_svg_path_gradient_reports_its_local_visual_fallback(self):
        from pptx_skill.document.svg import export_svg

        obj = shape(
            "rect-gradient",
            **{
                "paint.kind": "gradient",
                "paint.gradient.kind": "rect",
                "paint.gradient.fill_to_rect": {
                    "left": 0.5,
                    "top": 0.5,
                    "right": 0.5,
                    "bottom": 0.5,
                },
                "paint.gradient.stops": [
                    {"offset": 0, "color": "#244E9B"},
                    {"offset": 1, "color": "#F3DAAC"},
                ],
            },
        )
        report = export_svg(document([obj]), self.root / "gradient-svg")
        self.assertEqual(report["issues"][0]["status"], "raster_visual_fallback")
        self.assertEqual(report["issues"][0]["affected_objects"], ["rect-gradient"])
        root = etree.parse(report["pages"][0]["svg"])
        self.assertTrue(
            root.xpath(
                '//*[local-name()="image" and starts-with(@*[local-name()="href"],"data:image/png")]'
            )
        )

    def test_native_property_binding_checks_real_xml_target(self):
        source = document(properties={"document.default_language": "zh-CN"})
        source["native_bindings"] = [
            {
                "owner_id": "presentation",
                "property_id": "document.default_language",
                "part": "ppt/presentation.xml",
                "xpath": "/p:presentation/p:defaultTextStyle/a:defPPr/a:defRPr",
                "action": "set_attribute",
                "name": "lang",
                "value": "zh-CN",
            }
        ]
        path = self.root / "bound.pptx"
        report = export_pptx(source, path)
        with zipfile.ZipFile(path) as z:
            root = etree.fromstring(z.read("ppt/presentation.xml"))
            self.assertEqual(root.find(".//" + tag("a:defRPr")).get("lang"), "zh-CN")
        self.assertEqual(report["native_bindings"][0]["status"], "explicit_native_xml")
        source["native_bindings"][0]["xpath"] = "/p:presentation/nonexistent"
        with self.assertRaisesRegex(ValueError, "matched 0"):
            export_pptx(source, self.root / "bad-binding.pptx")

    def test_template_chart_edit_and_explicit_override_removal(self):
        chart = {
            "id": "chart",
            "kind": "chart",
            "properties": {
                "object.frame": {"left": 30, "top": 200, "width": 500, "height": 200},
                "chart.type": "barChart",
                "chart.categories": ["a", "b"],
                "chart.series": [{"name": "values", "values": [1, 2]}],
            },
        }
        text = {
            "id": "t",
            "kind": "text_box",
            "properties": {
                "object.frame": {"left": 30, "top": 30, "width": 500, "height": 100},
                "text.run.font.bold": True,
                "text.body.paragraphs": [
                    {
                        "content": [
                            {
                                "text": "title",
                                "properties": {"text.run.font.bold": True},
                            }
                        ]
                    }
                ],
            },
        }
        path = self.root / "template.pptx"
        export_pptx(document([chart, text]), path)
        doc = import_pptx(path)
        self.assertEqual([o.id for o in doc.pages[0].objects], ["chart", "t"])
        doc.pages[0].objects[0].properties["chart.series"][0]["values"] = [3, 4]
        doc.pages[0].objects[1].clear_overrides = ["text.run.font.bold"]
        out = self.root / "changed.pptx"
        export_pptx(doc, out)
        prs = Presentation(out)
        self.assertEqual(
            list(prs.slides[0].shapes[0].chart.series[0].values), [3.0, 4.0]
        )
        self.assertIsNone(
            prs.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.bold
        )

    def test_all_camera_light_material_and_underline_values(self):
        objects = []
        for key in (
            "three_d.camera.preset",
            "three_d.light_rig.preset",
            "three_d.material",
        ):
            for value in PROPERTIES[key]["allowed_values"]:
                objects.append(shape(key + "-" + value, **{key: value}))
        for value in PROPERTIES["text.run.underline.type"]["allowed_values"]:
            objects.append(
                {
                    "id": "underline-" + value,
                    "kind": "text_box",
                    "properties": {
                        "object.frame": {
                            "left": 30,
                            "top": 30,
                            "width": 300,
                            "height": 60,
                        },
                        "text.body.paragraphs": [
                            {
                                "content": [
                                    {
                                        "text": value,
                                        "properties": {
                                            "text.run.underline.type": value
                                        },
                                    }
                                ]
                            }
                        ],
                    },
                }
            )
        path = self.root / "more-enums.pptx"
        export_pptx(document(objects), path)
        prs = Presentation(path)
        self.assertEqual(len(prs.slides[0].shapes), 62 + 27 + 15 + 18)

    @unittest.skipUnless(
        __import__("shutil").which("soffice"), "SVG projection needs LibreOffice"
    )
    def test_svg_export_recovers_model_and_detects_visual_edits(self):
        from pptx_skill.document.svg import export_svg

        source = Document.model_validate(document([shape("stable-id")]))
        report = export_svg(source, self.root / "svg")
        path = Path(report["pages"][0]["svg"])
        drawing = etree.parse(str(path))
        slide = drawing.xpath(
            '//*[local-name()="g" and @class="Slide" and @id!="dummy-slide"]'
        )[0]
        self.assertFalse(
            any(
                n.get("visibility") == "hidden" or n.get("display") == "none"
                for n in [slide, *slide.iterancestors()]
            )
        )
        recovered, status = import_svg(path)
        self.assertEqual(recovered.pages[0].objects[0].id, "stable-id")
        self.assertTrue(status["lossless_source_recovery"])
        root = etree.parse(str(path))
        root.getroot().set("width", "999pt")
        path.write_bytes(etree.tostring(root))
        with self.assertRaisesRegex(ValueError, "stale"):
            import_svg(path)


if __name__ == "__main__":
    unittest.main(verbosity=2)
