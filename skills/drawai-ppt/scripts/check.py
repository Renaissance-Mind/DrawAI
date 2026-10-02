#!/usr/bin/env python3
"""Check the included toolkit with actual PPTX, SVG, bitmap and math files."""

from __future__ import annotations

import hashlib
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parent))

from PIL import Image
from pptx import Presentation
from pptx_skill import create_pptx, edit_pptx, inspect_pptx, validate_pptx_package


class ToolkitChecks(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="pptx-skill-check-")
        self.root = Path(self.directory.name)
        self.source = self.root / "source.pptx"
        self.frame = {"left": 80, "top": 180, "width": 400, "height": 160}
        create_pptx({"schema": "pptx.create_plan.v1", "slide_count": 2, "operations": [
            {"op": "add_text", "slide": 1, "text": "Original title", "name": "title",
             "frame": {"left": 80, "top": 40, "width": 900, "height": 80},
             "style": {"font_size": 32, "bold": True}},
            {"op": "add_table", "slide": 2, "frame": self.frame, "values": [["Input", "Output"], ["Plan", "PPTX"]]},
            {"op": "add_chart", "slide": 2, "frame": self.frame, "chart_type": "column",
             "data": {"categories": ["A", "B"], "series": [{"name": "S", "values": [10, 20]}]}},
        ]}, self.source)

    def tearDown(self):
        self.directory.cleanup()

    def edit(self, operations, name="edited.pptx"):
        output = self.root / name
        report = edit_pptx(self.source, {"schema": "pptx.edit_plan.v1", "operations": operations}, output, asset_root=self.root)
        self.assertEqual(report["status"], "ok")
        self.assertEqual(validate_pptx_package(output)["status"], "ok")
        return output, report

    def test_create_and_style_preserving_edit(self):
        inventory = inspect_pptx(self.source)
        self.assertEqual(len(inventory["slides"]), 2)
        original = self.source.read_bytes()
        title = inventory["slides"][0]["shapes"][0]
        output, _ = self.edit([{"op": "replace_text", "slide": 1, "shape_id": title["shape_id"],
                                "find": "Original", "text": "Updated"}])
        shape = Presentation(output).slides[0].shapes[0]
        self.assertEqual(shape.text, "Updated title")
        self.assertTrue(shape.text_frame.paragraphs[0].runs[0].font.bold)
        self.assertEqual(self.source.read_bytes(), original)

    def test_independent_chart_copy(self):
        chart = next(item for item in inspect_pptx(self.source)["slides"][1]["shapes"] if item["kind"] == "chart")
        output, _ = self.edit([
            {"op": "duplicate_slide", "slide": 2},
            {"op": "set_chart_data", "slide": 3, "shape_id": chart["shape_id"],
             "data": {"categories": ["A", "B"], "series": [{"name": "Copy", "values": [30, 40]}]}},
        ])
        deck = Presentation(output)
        self.assertEqual(list(deck.slides[1].shapes[1].chart.series[0].values), [10, 20])
        self.assertEqual(list(deck.slides[2].shapes[1].chart.series[0].values), [30, 40])

    def test_svg_paths_text_image_and_nonzero_viewbox(self):
        Image.new("RGB", (60, 40), "navy").save(self.root / "image.png")
        svg = self.root / "diagram.svg"
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="100 50 400 200">'
                       '<rect x="100" y="50" width="80" height="40" fill="#123456"/>'
                       '<path d="M 200 100 C 220 60 260 140 280 100" stroke="#456789" fill="none"/>'
                       '<text x="110" y="80" font-size="20">Editable text</text>'
                       '<image href="image.png" x="300" y="100" width="60" height="40"/></svg>')
        original = svg.read_bytes()
        output, report = self.edit([{"op": "add_svg", "slide": 1, "path": "diagram.svg", "frame": self.frame}])
        group = Presentation(output).slides[0].shapes[-1]
        self.assertEqual(group.shapes[0].left, 0)
        self.assertEqual(group.shapes[0].top, 0)
        self.assertTrue(any(shape.has_text_frame and shape.text == "Editable text" for shape in group.shapes))
        self.assertTrue(any(shape.shape_type == 13 for shape in group.shapes))
        self.assertEqual(svg.read_bytes(), original)
        self.assertTrue(report["operations"][0]["conversion"]["backend"])

    @unittest.skipUnless(shutil.which("pandoc"), "Office Math conversion requires Pandoc")
    def test_real_office_math(self):
        (self.root / "formula.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200">'
            '<g id="equation" data-pb-role="formula" data-pb-formula-latex="x^2+y^2" data-pb-formula-bbox="40 60 220 60">'
            '<text x="40" y="100" font-size="30">x²+y²</text></g></svg>')
        output, report = self.edit([{"op": "add_svg", "slide": 1, "path": "formula.svg", "frame": self.frame}])
        self.assertEqual(report["operations"][0]["conversion"]["formula_export"]["converted"], 1)
        self.assertTrue(any(item["kind"] == "office_math" for item in inspect_pptx(output)["slides"][0]["shapes"]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
