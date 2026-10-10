"""Exercise the common visual manifest and Flow bridge with actual files/API."""
import importlib.util
import json
import os
import shutil
from pathlib import Path
import socket
import subprocess
import sys
import threading
import time

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
import uvicorn

from drawai.workbench.api import create_app
from drawai.workbench.models import WorkbenchSettings

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills/drawai-image2ppt/scripts"
spec = importlib.util.spec_from_file_location("visual_deck", SCRIPTS / "visual_deck.py")
visual = importlib.util.module_from_spec(spec)
spec.loader.exec_module(visual)


@pytest.fixture
def actual_pages(tmp_path):
    # The inputs are rendered from actual native PowerPoint objects.
    sys.path.insert(0, str(SCRIPTS))
    from pptx_skill import render_pptx
    pages = []
    for index, name in enumerate(["B-first", "A-second"], 1):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(16), Inches(9)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(2)).text = name
        path = tmp_path / f"{name}.pptx"
        deck.save(path)
        preview = tmp_path / f"preview-{index}"
        render_pptx(path, preview)
        image = next(preview.glob("*.png"))
        pages.append({"id": f"page-{index}", "image": str(image), "title": name, "native_pptx": str(path)})
    input_spec = tmp_path / "visual-spec.json"
    input_spec.write_text(json.dumps({"title": "Flow integration", "aspect_ratio": "16:9", "pages": pages}))
    manifest = tmp_path / "visual-deck.json"
    visual.save_visual_deck(input_spec, manifest)
    return manifest, pages


def test_visual_manifest_verifies_actual_images_and_exposes_changed_bytes(actual_pages):
    manifest, pages = actual_pages
    result = visual.verify_visual_deck(manifest)
    assert [page["title"] for page in result["pages"]] == ["B-first", "A-second"]
    image = Path(pages[0]["image"])
    original = image.read_bytes()
    image.write_bytes(original + b"changed")
    with pytest.raises(ValueError, match="changed"):
        visual.verify_visual_deck(manifest)
    image.write_bytes(original)


def test_real_workbench_can_queue_without_launching_workers_and_bridge_exports_ordered_native_files(tmp_path, actual_pages):
    manifest, pages = actual_pages
    app = create_app(WorkbenchSettings(workspace=tmp_path / "workbench", default_config=ROOT / "configs/drawai/config.yaml"))
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    port = sock.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, log_level="error"))
    thread = threading.Thread(target=lambda: server.run(sockets=[sock]), daemon=True)
    thread.start()
    deadline = time.monotonic() + 15
    while not server.started and time.monotonic() < deadline:
        time.sleep(0.02)
    assert server.started
    env = {**os.environ, "DRAWAI_WORKBENCH_API_URL": f"http://127.0.0.1:{port}", "DRAWAI_PYTHON": sys.executable}
    cli = ROOT / "apps/workbench/templates/cli.mjs"

    def flow(arguments):
        args = tmp_path / "flow-args.json"
        args.write_text(json.dumps(arguments))
        result = subprocess.run(["node", str(cli), "drawai_flow", f"@{args}"], env=env, cwd=ROOT, capture_output=True, text=True, timeout=30)
        assert result.returncode == 0, result.stderr
        return json.loads(result.stdout)

    try:
        submitted = flow({"action": "submit", "visual_manifest": str(manifest)})
        batch_id = submitted["batch"]["batch_id"]
        cases = app.state.store.list_cases(batch_id)
        assert len(cases) == 2
        assert all(case.status == "queued" for case in cases)
        assert all(not app.state.store.list_stage_runs(case.case_id) for case in cases)
        assert [case.name[:5] for case in cases] == ["0001-", "0002-"]
        status = flow({"action": "status", "batch_id": batch_id})
        assert status["batch"]["status"] == "queued"
        # Register actual native results to test transport and assembly separately
        # from the GPU/model reconstruction execution.
        for case, page in zip(cases, pages):
            result_path = Path(case.run_root) / "exports" / "deck.pptx"
            result_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(page["native_pptx"], result_path)
            app.state.store.register_artifact(case.case_id, label="pptx", path=result_path, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
            app.state.store.update_case_status(case.case_id, status="completed", phase="export", stage="exported")
        with app.state.store._connect() as db:
            db.execute("UPDATE cases SET created_at = ? WHERE case_id = ?", ("2000-01-01", cases[1].case_id))
        output = tmp_path / "assembled.pptx"
        exported = flow({"action": "export", "batch_id": batch_id, "output_path": str(output)})
        assert exported["needs_validation"]
        deck = Presentation(output)
        assert [slide.shapes[0].text for slide in deck.slides] == ["B-first", "A-second"]
        sys.path.insert(0, str(SCRIPTS))
        from pptx_skill import validate_pptx_package
        assert validate_pptx_package(output)["status"] == "ok"
    finally:
        server.should_exit = True
        thread.join(timeout=10)
        sock.close()


def test_invalid_autostart_is_rejected_before_any_batch_is_created(tmp_path):
    app = create_app(WorkbenchSettings(workspace=tmp_path / "workbench", default_config=ROOT / "configs/drawai/config.yaml"))
    with TestClient(app) as client:
        response = client.post("/api/batches", json={"input_mode": "local_dir", "auto_start": "invalid"})
        assert response.status_code == 400
        assert app.state.store.list_batches() == []
