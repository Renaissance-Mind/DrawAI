from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from pydantic import ValidationError
import pytest

from drawai.cli import main
from drawai.pptx_editing import (
    EDIT_PLAN_SCHEMA, PptxEditError, edit_pptx, inspect_pptx, validate_pptx_package,
)
from drawai.tooling import drawai_tool_registry


@pytest.fixture
def template(tmp_path):
    image = tmp_path / "image.png"
    Image.new("RGB", (120, 80), "navy").save(image)
    replacement = tmp_path / "replacement.png"
    Image.new("RGB", (80, 120), "orange").save(replacement)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12), Inches(6.75)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(1))
    text.name = "heading"
    paragraph = text.text_frame.paragraphs[0]
    paragraph.space_after = Pt(12)
    run = paragraph.add_run()
    run.text = "Hello "
    run.font.name, run.font.size, run.font.bold = "Arial", Pt(24), True
    run.font.color.rgb = RGBColor(10, 20, 30)
    run = paragraph.add_run()
    run.text = "world!"
    run.font.italic = True
    run.font.color.rgb = RGBColor(60, 70, 80)
    run.hyperlink.address = "https://example.org/"
    image_shape = slide.shapes.add_picture(str(image), Inches(0.5), Inches(2), Inches(2), Inches(1.5))
    image_shape.crop_left = 0.1
    table = slide.shapes.add_table(2, 2, Inches(3), Inches(2), Inches(2), Inches(1))
    table.table.cell(0, 0).text = "Metric"
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Baseline", (1, 2))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(2), Inches(4), Inches(3), data)
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_textbox(Inches(0.5), Inches(5), Inches(4), Inches(0.5))
    nested.text = "Nested content"
    slide.notes_slide.notes_text_frame.text = "Speaker notes must survive."
    other = prs.slides.add_slide(prs.slide_layouts[6])
    other.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Untouched slide"
    source = tmp_path / "template.pptx"
    prs.save(source)
    return {"path": source, "text": text.shape_id, "image": image_shape.shape_id,
            "table": table.shape_id, "chart": chart.shape_id, "nested": nested.shape_id,
            "replacement": replacement, "image_path": image}


def plan(*operations):
    return {"schema": EDIT_PLAN_SCHEMA, "operations": list(operations)}


def get_shape(slide, shape_id):
    return next(shape for shape in slide.shapes if shape.shape_id == shape_id)


def test_inventory_exposes_nested_targets_and_layouts(template):
    result = inspect_pptx(template["path"])
    assert len(result["slides"]) == 2
    assert result["source_sha256"] == hashlib.sha256(template["path"].read_bytes()).hexdigest()
    assert {shape["kind"] for shape in result["slides"][0]["shapes"]} >= {"text", "table", "chart", "image", "group"}
    assert any(shape["shape_id"] == template["nested"] for shape in result["slides"][0]["shapes"])
    assert result["layouts"][0]["placeholders"]


def test_replacement_across_runs_preserves_formatting_and_source(template, tmp_path):
    source_bytes = template["path"].read_bytes()
    output = tmp_path / "edited.pptx"
    report = edit_pptx(template["path"], plan({"op": "replace_text", "slide": 1,
                                           "shape_id": template["text"], "find": "lo wor", "text": "p"}), output)
    prs = Presentation(output)
    paragraph = get_shape(prs.slides[0], template["text"]).text_frame.paragraphs[0]
    assert paragraph.text == "Helpld!"
    assert paragraph.runs[0].font.bold is True
    assert paragraph.runs[0].font.size == Pt(24)
    assert paragraph.runs[1].font.italic is True
    assert paragraph.runs[1].hyperlink.address == "https://example.org/"
    assert paragraph.space_after == Pt(12)
    assert template["path"].read_bytes() == source_bytes
    with zipfile.ZipFile(template["path"]) as before, zipfile.ZipFile(output) as after:
        for name in ("ppt/theme/theme1.xml", "ppt/slideMasters/slideMaster1.xml", "ppt/slides/slide2.xml"):
            assert before.read(name) == after.read(name)
    assert report["package_validation"]["status"] == "ok"


def test_set_text_keeps_paragraph_and_first_run_style(template, tmp_path):
    output = tmp_path / "text.pptx"
    edit_pptx(template["path"], plan({"op": "set_text", "slide": 1,
                                    "shape_id": template["text"], "text": "New title\nSecond line"},
                                   {"op": "replace_text", "slide": 1, "shape_id": template["nested"],
                                    "find": "Nested", "text": "Grouped"}), output)
    prs = Presentation(output)
    shape = get_shape(prs.slides[0], template["text"])
    assert shape.text == "New title\nSecond line"
    for paragraph in shape.text_frame.paragraphs:
        assert paragraph.runs[0].font.bold
        assert paragraph.runs[0].font.size == Pt(24)
        assert paragraph.space_after == Pt(12)
    assert "Grouped content" in json.dumps(inspect_pptx(output))


def test_native_image_table_chart_edits(template, tmp_path):
    output = tmp_path / "data.pptx"
    edit_pptx(template["path"], plan(
        {"op": "replace_image", "slide": 1, "shape_id": template["image"], "path": "replacement.png"},
        {"op": "set_table_cells", "slide": 1, "shape_id": template["table"],
         "cells": [{"row": 1, "column": 1, "text": "42", "style": {"bold": True}}]},
        {"op": "set_chart_data", "slide": 1, "shape_id": template["chart"],
         "data": {"categories": ["C", "D"], "series": [{"name": "Updated", "values": [10, 20]}]}},
    ), output)
    prs = Presentation(output)
    image = get_shape(prs.slides[0], template["image"])
    assert image.image.blob == template["replacement"].read_bytes()
    assert image.crop_left == pytest.approx(0.1)
    cell = get_shape(prs.slides[0], template["table"]).table.cell(1, 1)
    assert cell.text == "42" and cell.text_frame.paragraphs[0].runs[0].font.bold
    chart = get_shape(prs.slides[0], template["chart"]).chart
    assert chart.series[0].name == "Updated"
    assert list(chart.series[0].values) == [10, 20]
    assert [category.label for category in chart.plots[0].categories] == ["C", "D"]


def test_duplicate_slide_keeps_notes_and_clones_chart_workbook(template, tmp_path):
    output = tmp_path / "copied.pptx"
    edit_pptx(template["path"], plan(
        {"op": "duplicate_slide", "slide": 1},
        {"op": "set_chart_data", "slide": 3, "shape_id": template["chart"],
         "data": {"categories": ["A", "B"], "series": [{"name": "Copy", "values": [8, 9]}]}},
    ), output)
    prs = Presentation(output)
    original = get_shape(prs.slides[0], template["chart"]).chart
    copied = get_shape(prs.slides[2], template["chart"]).chart
    assert list(original.series[0].values) == [1, 2]
    assert list(copied.series[0].values) == [8, 9]
    assert original.part.partname != copied.part.partname
    assert original.part.chart_workbook.xlsx_part.partname != copied.part.chart_workbook.xlsx_part.partname
    assert prs.slides[2].notes_slide.notes_text_frame.text == "Speaker notes must survive."
    assert prs.slides[2].notes_slide.part.related_part("rId2") is prs.slides[2].part
    assert get_shape(prs.slides[2], template["text"]).text_frame.paragraphs[0].runs[1].hyperlink.address == "https://example.org/"


def test_insert_svg_reuses_native_converter_and_embeds_images(template, tmp_path):
    svg = tmp_path / "diagram.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200">'
                   '<rect x="5" y="5" width="120" height="80" rx="8" fill="#70a0e0"/>'
                   '<text x="15" y="40" font-size="20">Native SVG text</text>'
                   '<image href="image.png" x="150" y="20" width="120" height="80"/>'
                   '</svg>')
    source_svg = svg.read_bytes()
    output = tmp_path / "svg.pptx"
    result = edit_pptx(template["path"], plan(
        {"op": "add_svg", "slide": 1, "path": "diagram.svg",
         "frame": {"left": 500, "top": 20, "width": 400, "height": 200}},
        {"op": "add_svg", "slide": 1, "path": "diagram.svg",
         "frame": {"left": 500, "top": 400, "width": 200, "height": 100}},
    ), output)
    prs = Presentation(output)
    group = get_shape(prs.slides[0], result["operations"][0]["shape_id"])
    assert group.left == 500 * 9525 and group.top == 20 * 9525
    assert any(shape.has_text_frame and shape.text == "Native SVG text" for shape in group.shapes)
    assert any(shape.shape_type == 13 and shape.image.blob == template["image_path"].read_bytes() for shape in group.shapes)
    ids = [node.get("id") for node in prs.slides[0]._element.iter(qn("p:cNvPr"))]
    assert len(ids) == len(set(ids))
    assert svg.read_bytes() == source_svg
    assert not (tmp_path / "diagram.svg_to_ppt.svg").exists()


def test_add_layout_image_table_and_native_charts(template, tmp_path):
    layout = inspect_pptx(template["path"])["layouts"][0]["layout_id"]
    frame = {"left": 40, "top": 80, "width": 300, "height": 180}
    output = tmp_path / "added.pptx"
    result = edit_pptx(template["path"], plan(
        {"op": "add_slide", "layout_id": layout},
        {"op": "add_image", "slide": 3, "path": "image.png", "frame": frame, "fit": "cover"},
        {"op": "add_table", "slide": 3, "frame": frame, "values": [["Metric", "Value"], ["A", 42]]},
        {"op": "add_chart", "slide": 3, "frame": frame, "chart_type": "column",
         "data": {"categories": ["A"], "series": [{"name": "S", "values": [42]}]}},
        {"op": "add_chart", "slide": 3, "frame": frame, "chart_type": "scatter",
         "data": {"series": [{"name": "S", "points": [[1, 2], [3, 4]]}]}},
    ), output)
    prs = Presentation(output)
    assert len(prs.slides) == 3
    assert str(prs.slides[2].slide_layout.part.partname) == layout
    assert get_shape(prs.slides[2], result["operations"][2]["shape_id"]).table.cell(1, 1).text == "42"
    assert get_shape(prs.slides[2], result["operations"][3]["shape_id"]).has_chart
    assert get_shape(prs.slides[2], result["operations"][4]["shape_id"]).chart.chart_type == XL_CHART_TYPE.XY_SCATTER
    inspect_pptx(output)


@pytest.mark.parametrize("operation", [
    {"op": "set_text", "slide": 99, "shape_id": 1, "text": "bad"},
    {"op": "set_text", "slide": 1, "shape_id": 999, "text": "bad"},
    {"op": "replace_text", "slide": 1, "shape_id": 2, "find": "missing", "text": "bad"},
    {"op": "add_slide", "layout_id": "/missing/layout.xml"},
])
def test_bad_target_does_not_publish_output(template, tmp_path, operation):
    output = tmp_path / "invalid.pptx"
    with pytest.raises(PptxEditError):
        edit_pptx(template["path"], plan(operation), output)
    assert not output.exists()


def test_reject_stale_inspection_and_overwrite(template, tmp_path):
    with pytest.raises(PptxEditError, match="changed"):
        edit_pptx(template["path"], {**plan({"op": "duplicate_slide", "slide": 1}),
                                  "source_sha256": "0" * 64}, tmp_path / "stale.pptx")
    with pytest.raises(PptxEditError, match="new output"):
        edit_pptx(template["path"], plan({"op": "duplicate_slide", "slide": 1}), template["path"])
    with pytest.raises(ValidationError):
        edit_pptx(template["path"], plan({"op": "guess_operation"}), tmp_path / "unknown.pptx")


def test_geometry_and_explicit_styles(template, tmp_path):
    output = tmp_path / "style.pptx"
    edit_pptx(template["path"], plan(
        {"op": "set_geometry", "slide": 1, "shape_id": template["text"],
         "frame": {"left": 20, "top": 30, "width": 400, "height": 100}, "rotation": 5},
        {"op": "set_style", "slide": 1, "shape_id": template["text"],
         "fill": "#abcdef", "line_color": "#123456", "line_width": 2,
         "text_style": {"font_size": 20, "font_family": "Arial", "italic": True, "color": "#102030"}},
    ), output)
    shape = get_shape(Presentation(output).slides[0], template["text"])
    assert shape.left == 20 * 9525 and shape.rotation == 5
    assert str(shape.fill.fore_color.rgb) == "ABCDEF"
    assert shape.text_frame.paragraphs[0].runs[0].font.size == Pt(15)


def test_cli_and_agent_tool_share_actual_editor(template, tmp_path, capsys):
    inventory = tmp_path / "inventory.json"
    assert main(["tool", "ppt", "inspect", str(template["path"]), "--output", str(inventory)]) == 0
    assert json.loads(inventory.read_text())["schema"] == "drawai.pptx_inventory.v1"
    edit_plan = tmp_path / "plan.json"
    edit_plan.write_text(json.dumps(plan({"op": "duplicate_slide", "slide": 1})))
    output = tmp_path / "cli.pptx"
    assert main(["ppt", "edit", str(template["path"]), "--plan", str(edit_plan), "--output", str(output)]) == 0
    assert len(Presentation(output).slides) == 3
    assert main(["ppt", "validate", str(output)]) == 0
    assert drawai_tool_registry()["ppt"].tool_id == "ppt"
    assert main(["ppt", "edit", str(template["path"]), "--plan", str(edit_plan), "--output", str(output)]) == 2
    assert "never overwritten" in capsys.readouterr().err


def test_package_validator_reports_actual_broken_relationship(template, tmp_path):
    damaged = tmp_path / "damaged.pptx"
    with zipfile.ZipFile(template["path"]) as source, zipfile.ZipFile(damaged, "w") as target:
        for name in source.namelist():
            if name != "ppt/theme/theme1.xml":
                target.writestr(name, source.read(name))
    with pytest.raises(PptxEditError, match="Broken PPTX relationship"):
        validate_pptx_package(damaged)


def test_cli_json_destination_cannot_overwrite_a_source(template, capsys):
    source_bytes = template["path"].read_bytes()
    assert main(["ppt", "inspect", str(template["path"]), "--output", str(template["path"])]) == 2
    assert template["path"].read_bytes() == source_bytes
    assert "new .json filename" in capsys.readouterr().err


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Native formula conversion requires real Pandoc")
def test_svg_formula_insertion_keeps_native_office_math(template, tmp_path):
    svg = tmp_path / "formula.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 200">'
                   '<g id="equation" data-pb-role="formula" data-pb-formula-latex="x^2 + y^2" '
                   'data-pb-formula-bbox="40 60 220 60">'
                   '<text x="40" y="100" font-size="30">x² + y²</text></g></svg>')
    original = svg.read_bytes()
    output = tmp_path / "math.pptx"
    result = edit_pptx(template["path"], plan({"op": "add_svg", "slide": 1, "path": "formula.svg",
                                             "frame": {"left": 100, "top": 100, "width": 400, "height": 200}}), output)
    slide = Presentation(output).slides[0]
    group = get_shape(slide, result["operations"][0]["shape_id"])
    assert any(node.tag == "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
               for node in group._element.iter())
    assert result["operations"][0]["conversion"]["formula_export"]["converted"] == 1
    assert any(item["kind"] == "office_math" for item in inspect_pptx(output)["slides"][0]["shapes"])
    assert svg.read_bytes() == original
