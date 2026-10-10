from __future__ import annotations

import json
import os
from pathlib import Path
import shutil
import subprocess
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt
import pytest

from drawai.cli import main
from drawai.pptx_editing import CREATE_PLAN_SCHEMA, PptxEditError, create_pptx, edit_pptx, inspect_pptx


REPO = Path(__file__).resolve().parents[1]
LAUNCHER = REPO / "skills" / "drawai-ppt" / "scripts" / "ppt.py"


def create_plan():
    return {
        "schema": CREATE_PLAN_SCHEMA, "title": "Native PPTX workflow", "slide_count": 2,
        "slide_size": {"width": 1280, "height": 720},
        "operations": [
            {"op": "add_text", "slide": 1, "name": "title", "text": "Original title",
             "frame": {"left": 80, "top": 40, "width": 1000, "height": 80},
             "style": {"font_family": "Arial", "font_size": 40, "bold": True,
                       "alignment": "center", "vertical_alignment": "middle", "line_spacing": 44}},
            {"op": "add_shape", "slide": 1, "name": "node", "geometry": "round_rect",
             "frame": {"left": 80, "top": 180, "width": 400, "height": 200},
             "fill": "#E6EEF8", "line_color": "#537BAA", "line_width": 2,
             "text": "Editable node", "text_style": {"font_size": 28, "alignment": "center"}},
            {"op": "add_table", "slide": 2,
             "frame": {"left": 80, "top": 180, "width": 500, "height": 180},
             "values": [["Input", "Output"], ["Plan", "PPTX"]]},
            {"op": "add_chart", "slide": 2, "chart_type": "column",
             "frame": {"left": 650, "top": 180, "width": 500, "height": 300},
             "data": {"categories": ["A", "B"], "series": [{"name": "Series", "values": [10, 20]}]}},
        ],
    }


def test_create_native_objects_then_edit_generated_template(tmp_path):
    source = tmp_path / "created.pptx"
    report = create_pptx(create_plan(), source)
    prs = Presentation(source)
    assert report["schema"] == "drawai.pptx_create_report.v1"
    assert report["source_pptx"] is None
    assert report["slide_count"] == 2
    assert prs.core_properties.title == "Native PPTX workflow"
    assert prs.core_properties.author == "DrawAI"
    assert prs.slide_width == 1280 * 9525 and prs.slide_height == 720 * 9525
    text = prs.slides[0].shapes[0]
    assert text.name == "title" and text.text == "Original title"
    assert text.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
    paragraph = text.text_frame.paragraphs[0]
    assert paragraph.alignment == PP_ALIGN.CENTER
    assert paragraph.line_spacing == Pt(33)
    assert paragraph.runs[0].font.size == Pt(30)
    assert prs.slides[0].shapes[1].text == "Editable node"
    assert prs.slides[1].shapes[0].has_table and prs.slides[1].shapes[1].has_chart
    inventory = inspect_pptx(source)
    original_bytes = source.read_bytes()
    edited = tmp_path / "edited.pptx"
    edit_pptx(source, {"schema": "drawai.pptx_edit_plan.v1", "source_sha256": inventory["source_sha256"],
                       "operations": [{"op": "replace_text", "slide": 1, "shape_id": text.shape_id,
                                       "find": "Original", "text": "Updated"}]}, edited)
    reopened = Presentation(edited)
    assert reopened.slides[0].shapes[0].text == "Updated title"
    assert reopened.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size == Pt(30)
    assert source.read_bytes() == original_bytes


def test_skill_launcher_runs_outside_repo_and_after_relocation(tmp_path):
    source = tmp_path / "document.json"
    source.write_text(json.dumps({
        "schema": "presentation.document.v1", "id": "test-deck",
        "properties": {"document.page_size": {"width": 960, "height": 540}},
        "pages": [{"id": "overview", "objects": [{"id": "title", "kind": "text_box", "properties": {
            "object.frame": {"left": 60, "top": 30, "width": 700, "height": 80},
            "text.body.paragraphs": [{"content": [{"text": "Original title"}]}]
        }}]}]
    }))
    output = tmp_path / "deck.pptx"
    env = {"PATH": os.environ.get("PATH", os.defpath)}
    result = subprocess.run([sys.executable, "-I", str(LAUNCHER), "export", str(source),
                             "--pptx", str(output)], cwd=tmp_path, env=env,
                            check=True, capture_output=True, text=True)
    assert json.loads(result.stdout)["pptx"]["native_objects"] == 1
    relocated_root = tmp_path / "relocated"
    shutil.copytree(REPO / "skills" / "drawai-ppt", relocated_root,
                    ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
    relocated = relocated_root / "scripts" / "ppt.py"
    result = subprocess.run([sys.executable, "-I", str(relocated), "inspect", str(output)], cwd=tmp_path, env=env,
                            check=True, capture_output=True, text=True)
    inventory = json.loads(result.stdout)
    assert inventory["units"] == "pt"
    assert inventory["slides"][0]["shapes"][0]["text"] == "Original title"
    subprocess.run([sys.executable, "-I", str(relocated_root / "scripts" / "check.py")],
                   cwd=tmp_path, env=env, check=True, capture_output=True, text=True)


def test_create_cli_and_current_schema(tmp_path, capsys):
    assert main(["ppt", "schema", "--mode", "create"]) == 0
    schema = json.loads(capsys.readouterr().out)
    assert schema["properties"]["schema"]["const"] == CREATE_PLAN_SCHEMA
    assert "add_shape" in schema["properties"]["operations"]["items"]["discriminator"]["mapping"]
    path = tmp_path / "plan.json"
    path.write_text(json.dumps(create_plan()))
    assert main(["ppt", "create", "--plan", str(path), "--output", str(tmp_path / "cli.pptx"),
                 "--report", str(tmp_path / "report.json")]) == 0
    assert json.loads((tmp_path / "report.json").read_text())["status"] == "ok"


def test_invalid_shape_and_overwrite_fail_before_publication(tmp_path):
    plan = create_plan()
    plan["operations"][1]["geometry"] = "invented_geometry"
    output = tmp_path / "invalid.pptx"
    with pytest.raises(PptxEditError, match="Unsupported shape"):
        create_pptx(plan, output)
    assert not output.exists()
    create_pptx(create_plan(), output)
    original = output.read_bytes()
    with pytest.raises(PptxEditError, match="new output filename"):
        create_pptx(create_plan(), output)
    assert output.read_bytes() == original
