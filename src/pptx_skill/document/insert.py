"""Insert an SVG drawing through the document model, preserving native vectors."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from .model import load_document
from .importer import import_svg, import_pptx
from .export import export_pptx, patch_template_package
from .identity import fresh_identity, renew_identities, write_identity
from .xml import EMU, tag


def insert_svg(document, page_id, source, frame, *, fit="contain", asset_root=None):
    """Return (updated_document, conversion_report) with one native drawing group."""
    if fit not in ("contain", "stretch"):
        raise ValueError("SVG fit must be contain or stretch")
    if frame["width"] <= 0 or frame["height"] <= 0:
        raise ValueError("SVG frame must have positive point dimensions")
    doc = load_document(document)
    page_ids = [page.id for page in doc.pages]
    if page_id not in page_ids:
        raise ValueError(f"Unknown page ID: {page_id}")
    drawing, report = import_svg(source)
    projected = report.get("projection_page_id")
    index = next((i for i, page in enumerate(drawing.pages) if page.id == projected), 0)
    with tempfile.TemporaryDirectory() as tmpname:
        tmp = Path(tmpname)
        base = tmp / "base.pptx"
        vector = tmp / "vector.pptx"
        export_pptx(doc, base, asset_root=asset_root)
        export_pptx(drawing, vector, asset_root=Path(source).resolve().parent)
        target = Presentation(base)
        converted = Presentation(vector)
        slide = target.slides[page_ids.index(page_id)]
        source_slide = converted.slides[index]
        group = slide.shapes.add_group_shape()
        group_id = fresh_identity("svg-" + Path(source).stem)
        write_identity(group._element.find(".//" + tag("p:cNvPr")), group_id)
        ignored = {tag("p:nvGrpSpPr"), tag("p:grpSpPr"), tag("p:extLst")}
        elements = [
            deepcopy(node)
            for node in source_slide._element.spTree
            if node.tag not in ignored
        ]
        next_id = (
            max(int(node.get("id")) for node in slide._element.iter(tag("p:cNvPr"))) + 1
        )
        ids = {}
        for element in elements:
            renew_identities(element)
            for node in element.iter(tag("p:cNvPr")):
                ids[node.get("id")] = str(next_id)
                node.set("id", str(next_id))
                next_id += 1
        relationships = {}
        rns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        for element in elements:
            for node in element.iter():
                if (
                    node.tag in (tag("a:stCxn"), tag("a:endCxn"))
                    and node.get("id") in ids
                ):
                    node.set("id", ids[node.get("id")])
                for attr, rid in list(node.attrib.items()):
                    if not attr.startswith(rns):
                        continue
                    if rid not in relationships:
                        rel = source_slide.part.rels[rid]
                        if rel.is_external:
                            relationships[rid] = slide.part.relate_to(
                                rel.target_ref, rel.reltype, is_external=True
                            )
                        elif rel.reltype == RT.IMAGE:
                            _, relationships[rid] = slide.part.get_or_add_image_part(
                                BytesIO(rel.target_part.blob)
                            )
                        else:
                            raise ValueError(
                                f"SVG insertion needs standalone drawing resources; dedicated part {rel.reltype} requires document editing"
                            )
                    node.set(attr, relationships[rid])
            group._element.insert_element_before(element, "p:extLst")
        left, top, width, height = [
            round(frame[k] * EMU) for k in ("left", "top", "width", "height")
        ]
        vw, vh = converted.slide_width, converted.slide_height
        if fit == "contain":
            scale = min(width / vw, height / vh)
            w, h = round(vw * scale), round(vh * scale)
            left += (width - w) // 2
            top += (height - h) // 2
            width, height = w, h
        transform = group._element.grpSpPr.get_or_add_xfrm()
        transform.off.x, transform.off.y = left, top
        transform.ext.cx, transform.ext.cy = width, height
        transform.chOff.x, transform.chOff.y = 0, 0
        transform.chExt.cx, transform.chExt.cy = vw, vh
        saved = BytesIO()
        target.save(saved)
        out = tmp / "inserted.pptx"
        out.write_bytes(patch_template_package(base.read_bytes(), saved.getvalue()))
        updated = import_pptx(out)
    return updated, {
        **report,
        "inserted_object_id": group_id,
        "page_id": page_id,
        "fit": fit,
    }
