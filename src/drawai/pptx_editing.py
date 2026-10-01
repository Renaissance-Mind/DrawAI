"""Edit real PPTX templates and insert DrawAI's native SVG objects."""

from __future__ import annotations

from copy import deepcopy
import hashlib
from io import BytesIO
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
import shutil
import subprocess
import tempfile
from typing import Annotated, Any, Literal
from urllib.parse import unquote
import zipfile

from lxml import etree
from pydantic import BaseModel, ConfigDict, Field
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, _Relationship
from pptx.opc.packuri import PackURI
from pptx.util import Pt


EMU_PER_PX = 9525
EDIT_PLAN_SCHEMA = "drawai.pptx_edit_plan.v1"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


class PptxEditError(ValueError):
    """A plan cannot be applied faithfully to the requested presentation."""


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)


class Frame(StrictModel):
    left: float
    top: float
    width: float = Field(gt=0)
    height: float = Field(gt=0)


class TextStyle(StrictModel):
    font_family: str | None = None
    font_size: float | None = Field(default=None, gt=0)
    bold: bool | None = None
    italic: bool | None = None
    underline: bool | None = None
    color: str | None = Field(default=None, pattern=r"^#?[0-9a-fA-F]{6}$")


class TargetOperation(StrictModel):
    slide: int = Field(ge=1)
    shape_id: int = Field(ge=1)


class SetText(TargetOperation):
    op: Literal["set_text"]
    text: str
    style: TextStyle | None = None


class ReplaceText(TargetOperation):
    op: Literal["replace_text"]
    find: str = Field(min_length=1)
    text: str


class ReplaceImage(TargetOperation):
    op: Literal["replace_image"]
    path: str
    fit: Literal["preserve", "cover"] = "preserve"


class CellEdit(StrictModel):
    row: int = Field(ge=0)
    column: int = Field(ge=0)
    text: str
    style: TextStyle | None = None


class SetTableCells(TargetOperation):
    op: Literal["set_table_cells"]
    cells: list[CellEdit] = Field(min_length=1)


class ChartSeries(StrictModel):
    name: str
    values: list[float | None] | None = None
    points: list[list[float]] | None = None


class ChartData(StrictModel):
    categories: list[str] | None = None
    series: list[ChartSeries] = Field(min_length=1)


class SetChartData(TargetOperation):
    op: Literal["set_chart_data"]
    data: ChartData


class SetGeometry(TargetOperation):
    op: Literal["set_geometry"]
    frame: Frame
    rotation: float | None = None


class SetStyle(TargetOperation):
    op: Literal["set_style"]
    fill: str | None = Field(default=None, pattern=r"^(none|#?[0-9a-fA-F]{6})$")
    line_color: str | None = Field(default=None, pattern=r"^(none|#?[0-9a-fA-F]{6})$")
    line_width: float | None = Field(default=None, ge=0)
    text_style: TextStyle | None = None


class AddSvg(StrictModel):
    op: Literal["add_svg"]
    slide: int = Field(ge=1)
    path: str
    frame: Frame
    fit: Literal["contain", "stretch"] = "contain"
    name: str = "DrawAI SVG"


class AddImage(StrictModel):
    op: Literal["add_image"]
    slide: int = Field(ge=1)
    path: str
    frame: Frame
    fit: Literal["cover", "contain"] = "contain"


class AddTable(StrictModel):
    op: Literal["add_table"]
    slide: int = Field(ge=1)
    frame: Frame
    values: list[list[str | float | int]] = Field(min_length=1)
    style: TextStyle | None = None


class AddChart(StrictModel):
    op: Literal["add_chart"]
    slide: int = Field(ge=1)
    frame: Frame
    chart_type: Literal["column", "bar", "line", "area", "pie", "doughnut", "scatter", "bubble"]
    data: ChartData


class AddSlide(StrictModel):
    op: Literal["add_slide"]
    layout_id: str


class DuplicateSlide(StrictModel):
    op: Literal["duplicate_slide"]
    slide: int = Field(ge=1)


Operation = Annotated[
    SetText | ReplaceText | ReplaceImage | SetTableCells | SetChartData | SetGeometry
    | SetStyle | AddSvg | AddImage | AddTable | AddChart | AddSlide | DuplicateSlide,
    Field(discriminator="op"),
]


class PptxEditPlan(StrictModel):
    schema_: Literal["drawai.pptx_edit_plan.v1"] = Field(alias="schema")
    source_sha256: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")
    operations: list[Operation] = Field(min_length=1)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _source(path: str | Path) -> Path:
    source = Path(path).expanduser().resolve(strict=True)
    validate_pptx_package(source)
    return source


def _walk_shapes(shapes: Any):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)


def _frame(shape: Any) -> dict[str, float]:
    return {key: round(float(getattr(shape, key)) / EMU_PER_PX, 4)
            for key in ("left", "top", "width", "height")}


def _text_inventory(frame: Any) -> list[dict[str, Any]]:
    return [{"text": paragraph.text, "alignment": str(paragraph.alignment),
             "runs": [{"text": run.text, "font_family": run.font.name,
                       "font_size": float(run.font.size) / EMU_PER_PX if run.font.size else None,
                       "bold": run.font.bold, "italic": run.font.italic}
                      for run in paragraph.runs]}
            for paragraph in frame.paragraphs]


def inspect_pptx(path: str | Path) -> dict[str, Any]:
    """Return slide-local shape IDs, editable content, layouts and geometry."""
    source = _source(path)
    prs = Presentation(str(source))
    slides = []
    for number, slide in enumerate(prs.slides, 1):
        shapes = []
        parents = {child.shape_id: group.shape_id for group in _walk_shapes(slide.shapes)
                   if group.shape_type == MSO_SHAPE_TYPE.GROUP for child in group.shapes}
        for shape in _walk_shapes(slide.shapes):
            item = {"shape_id": shape.shape_id, "name": shape.name,
                    "type": str(shape.shape_type), "frame": _frame(shape),
                    "rotation": shape.rotation, "is_placeholder": shape.is_placeholder,
                    "parent_shape_id": parents.get(shape.shape_id),
                    "coordinate_space": "group" if shape.shape_id in parents else "slide"}
            if shape.is_placeholder:
                item["placeholder"] = {"idx": shape.placeholder_format.idx,
                                       "type": str(shape.placeholder_format.type)}
            if shape.has_text_frame:
                item.update(kind="text", text=shape.text,
                            paragraphs=_text_inventory(shape.text_frame))
            elif shape.has_table:
                item.update(kind="table", values=[[cell.text for cell in row.cells]
                                                   for row in shape.table.rows])
            elif shape.has_chart:
                item.update(kind="chart", chart_type=str(shape.chart.chart_type),
                            series=[{"name": series.name, "values": list(series.values)}
                                    for series in shape.chart.series])
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                item.update(kind="image", content_type=shape.image.content_type,
                            image_size=list(shape.image.size),
                            crop={key: getattr(shape, f"crop_{key}")
                                  for key in ("left", "top", "right", "bottom")})
            else:
                item["kind"] = "group" if shape.shape_type == MSO_SHAPE_TYPE.GROUP else "shape"
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    transform = shape._element.grpSpPr.xfrm
                    if transform is not None:
                        item["child_frame"] = {
                            "left": transform.chOff.x / EMU_PER_PX,
                            "top": transform.chOff.y / EMU_PER_PX,
                            "width": transform.chExt.cx / EMU_PER_PX,
                            "height": transform.chExt.cy / EMU_PER_PX,
                        }
            shapes.append(item)
        # python-pptx omits drawing entries inside markup compatibility choices.
        # Inventory them explicitly rather than hiding preserved Office Math.
        known_ids = {item["shape_id"] for item in shapes}
        containers = {qn("p:sp"), qn("p:pic"), qn("p:grpSp"), qn("p:graphicFrame"), qn("p:cxnSp")}
        for properties in slide._element.spTree.iter(qn("p:cNvPr")):
            shape_id = int(properties.get("id"))
            if shape_id in known_ids:
                continue
            owner = properties.getparent()
            while owner is not None and owner.tag not in containers:
                owner = owner.getparent()
            if owner is None:
                continue
            transform = owner.find("./" + qn("p:spPr") + "/" + qn("a:xfrm"))
            frame = None
            if transform is not None:
                off, ext = transform.find(qn("a:off")), transform.find(qn("a:ext"))
                if off is not None and ext is not None:
                    frame = {"left": int(off.get("x")) / EMU_PER_PX, "top": int(off.get("y")) / EMU_PER_PX,
                             "width": int(ext.get("cx")) / EMU_PER_PX, "height": int(ext.get("cy")) / EMU_PER_PX}
            parent_group = next((node for node in owner.iterancestors() if node.tag == qn("p:grpSp")), None)
            parent_id = None if parent_group is None else int(parent_group.find("./" + qn("p:nvGrpSpPr") + "/" + qn("p:cNvPr")).get("id"))
            math_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/math"
            is_math = any(node.tag == f"{{{math_namespace}}}oMath" for node in owner.iter())
            shapes.append({"shape_id": shape_id, "name": properties.get("name", ""),
                           "kind": "office_math" if is_math else "preserved_drawing",
                           "frame": frame, "parent_shape_id": parent_id,
                           "coordinate_space": "group" if parent_id is not None else "slide",
                           "editable_via_text_api": False})
            known_ids.add(shape_id)
        slides.append({"slide": number, "part": str(slide.part.partname),
                       "layout_id": str(slide.slide_layout.part.partname),
                       "shapes": shapes})
    layouts = [{"layout_id": str(layout.part.partname), "name": layout.name,
                "master_id": str(master.part.partname),
                "placeholders": [{"idx": shape.placeholder_format.idx,
                                   "name": shape.name, "frame": _frame(shape),
                                   "type": str(shape.placeholder_format.type)}
                                  for shape in layout.placeholders]}
               for master in prs.slide_masters for layout in master.slide_layouts]
    return {"schema": "drawai.pptx_inventory.v1", "source_pptx": str(source),
            "source_sha256": _sha256(source), "units": "px at 96 DPI",
            "slide_size": {"width": prs.slide_width / EMU_PER_PX,
                           "height": prs.slide_height / EMU_PER_PX},
            "layouts": layouts, "slides": slides}


def validate_pptx_package(path: str | Path) -> dict[str, Any]:
    """Check actual package XML and internal relationship destinations."""
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        if len(names) != len(set(names)):
            raise PptxEditError("PPTX contains duplicate package entries")
        required = {"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"}
        if not required.issubset(names):
            raise PptxEditError("File is not a PPTX presentation package")
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        xml_files = {name: etree.fromstring(package.read(name), parser)
                     for name in names if name.endswith((".xml", ".rels"))}
        relationship_ids = {}
        for name, root in xml_files.items():
            if not name.endswith(".rels"):
                continue
            base = posixpath.dirname(posixpath.dirname(name))
            ids = [rel.get("Id") for rel in root]
            if len(ids) != len(set(ids)) or None in ids:
                raise PptxEditError(f"Invalid relationship IDs in {name}")
            relationship_ids[name] = set(ids)
            for rel in root:
                if rel.get("TargetMode") == "External":
                    continue
                target = unquote(rel.get("Target", "").split("#", 1)[0])
                resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                if resolved not in names:
                    raise PptxEditError(f"Broken PPTX relationship: {name} -> {target}")
        for name, root in xml_files.items():
            if name.endswith(".rels"):
                continue
            path = PurePosixPath(name)
            rels_name = str(path.parent / "_rels" / (path.name + ".rels"))
            ids = relationship_ids.get(rels_name, set())
            for node in root.iter():
                for attr, value in node.attrib.items():
                    if attr.startswith("{" + _R_NS + "}") and value not in ids:
                        raise PptxEditError(f"Unknown resource relationship in {name}: {value}")
    return {"status": "ok", "package_entries": len(names), "xml_parts": len(xml_files)}


def _slide(prs: Any, number: int):
    if number > len(prs.slides):
        raise PptxEditError(f"Slide {number} does not exist (deck has {len(prs.slides)} slides)")
    return prs.slides[number - 1]


def _shape(slide: Any, shape_id: int):
    for shape in _walk_shapes(slide.shapes):
        if shape.shape_id == shape_id:
            return shape
    raise PptxEditError(f"Shape {shape_id} does not exist on the selected slide")


def _style_text(frame: Any, style: TextStyle | None) -> None:
    if style is None:
        return
    for paragraph in frame.paragraphs:
        fonts = [paragraph.font, *(run.font for run in paragraph.runs)]
        for font in fonts:
            if style.font_family is not None:
                font.name = style.font_family
            if style.font_size is not None:
                font.size = Pt(style.font_size * 0.75)
            for key in ("bold", "italic", "underline"):
                if getattr(style, key) is not None:
                    setattr(font, key, getattr(style, key))
            if style.color is not None:
                font.color.rgb = RGBColor.from_string(style.color.lstrip("#"))


def _set_text(frame: Any, text: str, style: TextStyle | None = None) -> None:
    """Keep paragraph properties and each template paragraph's first run style."""
    originals = [deepcopy(paragraph._p) for paragraph in frame.paragraphs]
    for paragraph in list(frame._txBody.findall(qn("a:p"))):
        frame._txBody.remove(paragraph)
    for index, line in enumerate(text.split("\n")):
        paragraph = deepcopy(originals[min(index, len(originals) - 1)])
        run = paragraph.find(qn("a:r"))
        run_properties = deepcopy(run.find(qn("a:rPr"))) if run is not None and run.find(qn("a:rPr")) is not None else None
        for child in list(paragraph):
            if child.tag not in {qn("a:pPr"), qn("a:endParaRPr")}:
                paragraph.remove(child)
        run = OxmlElement("a:r")
        if run_properties is not None:
            run.append(run_properties)
        value = OxmlElement("a:t")
        value.text = line
        run.append(value)
        paragraph.insert_element_before(run, "a:endParaRPr")
        frame._txBody.append(paragraph)
    _style_text(frame, style)


def _replace_text(frame: Any, find: str, text: str) -> int:
    if any(char in find + text for char in ("\n", "\v")):
        raise PptxEditError("replace_text operates within paragraphs; use set_text for multiline content")
    count = 0
    for paragraph in frame.paragraphs:
        # Breaks prevent a match from crossing a visual line boundary.
        segments = []
        offset = 0
        for child in paragraph._p:
            value = child.find(qn("a:t"))
            content = value.text or "" if value is not None else "\v" if child.tag == qn("a:br") else ""
            if content:
                segments.append((offset, offset + len(content), value, content))
                offset += len(content)
        original = "".join(segment[3] for segment in segments)
        matches = list(re.finditer(re.escape(find), original))
        for match in reversed(matches):
            inserted = False
            for start, end, value, _ in segments:
                if end <= match.start() or start >= match.end():
                    continue
                if value is None:
                    raise PptxEditError("Replacement crosses a text break")
                current = value.text or ""
                left = max(match.start() - start, 0)
                right = min(match.end() - start, end - start)
                value.text = current[:left] + (text if not inserted else "") + current[right:]
                inserted = True
            count += 1
    if not count:
        raise PptxEditError(f"Text to replace was not found: {find!r}")
    return count


def _asset(path: str, root: Path) -> Path:
    asset = Path(path).expanduser()
    return (asset if asset.is_absolute() else root / asset).resolve(strict=True)


def _dimensions(frame: Frame) -> tuple[int, int, int, int]:
    return tuple(round(getattr(frame, key) * EMU_PER_PX) for key in ("left", "top", "width", "height"))


def _chart_data(data: ChartData, chart_type: Any):
    scatter = chart_type in {XL_CHART_TYPE.XY_SCATTER, XL_CHART_TYPE.XY_SCATTER_LINES,
                             XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
                             XL_CHART_TYPE.XY_SCATTER_SMOOTH, XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS}
    bubble = chart_type in {XL_CHART_TYPE.BUBBLE, XL_CHART_TYPE.BUBBLE_THREE_D_EFFECT}
    result = BubbleChartData() if bubble else XyChartData() if scatter else CategoryChartData()
    if not scatter and not bubble:
        if not data.categories:
            raise PptxEditError("A category chart requires categories")
        result.categories = data.categories
    for series in data.series:
        if scatter or bubble:
            if not series.points or series.values is not None or data.categories is not None:
                raise PptxEditError("Scatter/bubble data requires points, without categories or values")
            item = result.add_series(series.name)
            for point in series.points:
                if len(point) != (3 if bubble else 2):
                    raise PptxEditError("Scatter points require [x,y]; bubble points require [x,y,size]")
                item.add_data_point(*point)
        else:
            if series.values is None or len(series.values) != len(data.categories) or series.points is not None:
                raise PptxEditError("Each chart series must have one value per category")
            result.add_series(series.name, series.values)
    return result


def _copy_slide(prs: Any, source: Any):
    """Copy a slide, cloning mutable chart/workbook/notes parts with their rIds."""
    target = prs.slides.add_slide(source.slide_layout)
    target.part._element = deepcopy(source._element)
    # Slide wraps its element separately from SlidePart.
    target._element = target.part._element
    # add_slide() populates placeholders through a cached SlideShapes facade.
    # Rebind that facade to the copied shape tree before further edits.
    target.__dict__.pop("shapes", None)
    target.__dict__.pop("placeholders", None)
    target.part.rels._rels.clear()
    mapping = {source.part: target.part}
    reserved = {str(part.partname) for part in prs.part.package.iter_parts()}
    shared = {RT.SLIDE_LAYOUT, RT.SLIDE_MASTER, RT.THEME, RT.IMAGE, RT.SLIDE}

    def clone(part: Part):
        if part in mapping:
            return mapping[part]
        path = PurePosixPath(str(part.partname))
        number = 1
        candidate = str(path.with_name(f"{path.stem}_drawai{number}{path.suffix}"))
        while candidate in reserved:
            number += 1
            candidate = str(path.with_name(f"{path.stem}_drawai{number}{path.suffix}"))
        reserved.add(candidate)
        copied = type(part).load(PackURI(candidate), part.content_type, prs.part.package, part.blob)
        mapping[part] = copied
        relationships(part, copied)
        return copied

    def relationships(original: Part, copied: Part):
        for rel in original.rels.values():
            if rel.is_external:
                destination = rel.target_ref
            elif rel.target_part in mapping:
                destination = mapping[rel.target_part]
            elif rel.reltype in shared:
                destination = rel.target_part
            else:
                destination = clone(rel.target_part)
            copied.rels._rels[rel.rId] = _Relationship(
                copied.partname.baseURI, rel.rId, rel.reltype,
                "External" if rel.is_external else "Internal", destination,
            )

    relationships(source.part, target.part)
    return target


def _insert_svg(prs: Any, slide: Any, operation: AddSvg, root: Path, temp: Path) -> dict[str, Any]:
    from .svg_to_ppt import SvgToPptCompiler

    svg = _asset(operation.path, root)
    # Preparation writes beside its input. Work with a copy and resolve asset
    # hrefs against the original SVG so user source files remain untouched.
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    svg_root = etree.fromstring(svg.read_bytes(), parser)
    for image in svg_root.iter("{http://www.w3.org/2000/svg}image"):
        for attr in ("href", "{http://www.w3.org/1999/xlink}href"):
            href = image.get(attr)
            if href and href.startswith(("http:", "https:")):
                raise PptxEditError("SVG insertion requires local image assets")
            if href and not href.startswith(("data:", "http:", "https:")):
                image.set(attr, str(_asset(href, svg.parent)))
    svg_copy = temp / "insert.svg"
    svg_copy.write_bytes(etree.tostring(svg_root))
    converted_path = temp / "insert.pptx"
    report = SvgToPptCompiler().compile(svg_copy, converted_path, check_full_slide_editability=False)
    if report.get("status") != "ok":
        raise PptxEditError(f"SVG conversion failed: {report}")
    converted = Presentation(str(converted_path))
    source_slide = converted.slides[0]
    group = slide.shapes.add_group_shape()
    group.name = operation.name
    next_id = max(shape.shape_id for shape in _walk_shapes(slide.shapes)) + 1
    id_map = {}
    # The public shape collection omits mc:AlternateContent, including native
    # Office Math. Copy the actual drawing-tree entries so those objects survive.
    utility_tags = {qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst")}
    elements = [deepcopy(element) for element in source_slide._element.spTree
                if element.tag not in utility_tags]
    for element in elements:
        for properties in element.iter(qn("p:cNvPr")):
            old_id = properties.get("id")
            id_map[old_id] = str(next_id)
            properties.set("id", str(next_id))
            next_id += 1
    rel_map = {}
    for element in elements:
        for node in element.iter():
            if node.tag in {qn("a:stCxn"), qn("a:endCxn")} and node.get("id") in id_map:
                node.set("id", id_map[node.get("id")])
            for attr, old_id in list(node.attrib.items()):
                if not attr.startswith("{" + _R_NS + "}"):
                    continue
                if old_id not in rel_map:
                    rel = source_slide.part.rels[old_id]
                    if rel.is_external:
                        rel_map[old_id] = slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                    elif rel.reltype == RT.IMAGE:
                        _, rel_map[old_id] = slide.part.get_or_add_image_part(BytesIO(rel.target_part.blob))
                    else:
                        raise PptxEditError(f"Unsupported SVG resource relationship: {rel.reltype}")
                node.set(attr, rel_map[old_id])
        group._element.insert_element_before(element, "p:extLst")
    left, top, width, height = _dimensions(operation.frame)
    viewport_w, viewport_h = converted.slide_width, converted.slide_height
    if operation.fit == "contain":
        scale = min(width / viewport_w, height / viewport_h)
        fitted_w, fitted_h = round(viewport_w * scale), round(viewport_h * scale)
        left += (width - fitted_w) // 2
        top += (height - fitted_h) // 2
        width, height = fitted_w, fitted_h
    transform = group._element.grpSpPr.get_or_add_xfrm()
    transform.off.x, transform.off.y = left, top
    transform.ext.cx, transform.ext.cy = width, height
    transform.chOff.x, transform.chOff.y = 0, 0
    transform.chExt.cx, transform.chExt.cy = viewport_w, viewport_h
    formula = report.get("formula_export", {})
    return {"shape_id": group.shape_id, "conversion": {
        "backend": report["backend"], "editable_surface": report["editable_surface"],
        "formula_export": {key: formula.get(key) for key in ("status", "count", "converted", "fallback", "conversion_issues")},
    }}


def _apply_operation(prs: Any, operation: Any, root: Path, temp: Path) -> dict[str, Any]:
    if isinstance(operation, AddSlide):
        layouts = {str(layout.part.partname): layout for master in prs.slide_masters for layout in master.slide_layouts}
        if operation.layout_id not in layouts:
            raise PptxEditError(f"Unknown layout_id: {operation.layout_id}")
        prs.slides.add_slide(layouts[operation.layout_id])
        return {"slide": len(prs.slides)}
    slide = _slide(prs, operation.slide)
    if isinstance(operation, DuplicateSlide):
        _copy_slide(prs, slide)
        return {"slide": len(prs.slides), "source_slide": operation.slide}
    result = {"slide": operation.slide}
    if isinstance(operation, AddSvg):
        result.update(_insert_svg(prs, slide, operation, root, temp))
        return result
    if isinstance(operation, (AddImage, AddTable, AddChart)):
        frame = _dimensions(operation.frame)
        if isinstance(operation, AddImage):
            shape = slide.shapes.add_picture(str(_asset(operation.path, root)), *frame)
            if operation.fit == "cover":
                _cover_image(shape)
            else:
                image_w, image_h = shape.image.size
                scale = min(frame[2] / image_w, frame[3] / image_h)
                shape.width, shape.height = round(image_w * scale), round(image_h * scale)
                shape.left = frame[0] + (frame[2] - shape.width) // 2
                shape.top = frame[1] + (frame[3] - shape.height) // 2
        elif isinstance(operation, AddTable):
            columns = len(operation.values[0])
            if not columns or any(len(row) != columns for row in operation.values):
                raise PptxEditError("Table values must be a nonempty rectangular matrix")
            shape = slide.shapes.add_table(len(operation.values), columns, *frame)
            for row, values in enumerate(operation.values):
                for column, value in enumerate(values):
                    _set_text(shape.table.cell(row, column).text_frame, str(value), operation.style)
        else:
            types = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                     "line": XL_CHART_TYPE.LINE, "area": XL_CHART_TYPE.AREA,
                     "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
                     "scatter": XL_CHART_TYPE.XY_SCATTER, "bubble": XL_CHART_TYPE.BUBBLE}
            kind = types[operation.chart_type]
            shape = slide.shapes.add_chart(kind, *frame, _chart_data(operation.data, kind))
        result["shape_id"] = shape.shape_id
        return result
    shape = _shape(slide, operation.shape_id)
    result["shape_id"] = shape.shape_id
    if isinstance(operation, (SetText, ReplaceText)):
        if not shape.has_text_frame:
            raise PptxEditError(f"Shape {shape.shape_id} has no editable text frame")
        if isinstance(operation, SetText):
            _set_text(shape.text_frame, operation.text, operation.style)
        else:
            result["replacements"] = _replace_text(shape.text_frame, operation.find, operation.text)
    elif isinstance(operation, ReplaceImage):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            raise PptxEditError("replace_image requires a picture object")
        _, rel_id = slide.part.get_or_add_image_part(str(_asset(operation.path, root)))
        blip = shape._element.xpath(".//a:blip")[0]
        blip.set(qn("r:embed"), rel_id)
        blip.attrib.pop(qn("r:link"), None)
        if operation.fit == "cover":
            _cover_image(shape)
    elif isinstance(operation, SetTableCells):
        if not shape.has_table:
            raise PptxEditError("set_table_cells requires a native table")
        for cell in operation.cells:
            if cell.row >= len(shape.table.rows) or cell.column >= len(shape.table.columns):
                raise PptxEditError("Table cell index is outside the table")
            _set_text(shape.table.cell(cell.row, cell.column).text_frame, cell.text, cell.style)
    elif isinstance(operation, SetChartData):
        if not shape.has_chart:
            raise PptxEditError("set_chart_data requires a native chart")
        shape.chart.replace_data(_chart_data(operation.data, shape.chart.chart_type))
    elif isinstance(operation, SetGeometry):
        shape.left, shape.top, shape.width, shape.height = _dimensions(operation.frame)
        if operation.rotation is not None:
            shape.rotation = operation.rotation
    elif isinstance(operation, SetStyle):
        if operation.fill is not None:
            if not hasattr(shape, "fill"):
                raise PptxEditError("This object has no shape fill; edit native table cells or chart regions directly")
            if operation.fill == "none":
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(operation.fill.lstrip("#"))
        if operation.line_color is not None:
            if operation.line_color == "none":
                shape.line.fill.background()
            else:
                shape.line.color.rgb = RGBColor.from_string(operation.line_color.lstrip("#"))
        if operation.line_width is not None:
            shape.line.width = Pt(operation.line_width * 0.75)
        if operation.text_style is not None:
            if not shape.has_text_frame:
                raise PptxEditError("text_style requires an editable text frame")
            _style_text(shape.text_frame, operation.text_style)
    return result


def _cover_image(shape: Any) -> None:
    image_w, image_h = shape.image.size
    image_ratio = image_w / image_h
    frame_ratio = shape.width / shape.height
    shape.crop_left = shape.crop_right = max(0, (1 - frame_ratio / image_ratio) / 2)
    shape.crop_top = shape.crop_bottom = max(0, (1 - image_ratio / frame_ratio) / 2)


def edit_pptx(source_path: str | Path, plan: dict[str, Any] | PptxEditPlan,
              output_path: str | Path, *, asset_root: str | Path | None = None) -> dict[str, Any]:
    """Apply a strict plan, publishing a new file only after package validation."""
    parsed = plan if isinstance(plan, PptxEditPlan) else PptxEditPlan.model_validate(plan)
    source = _source(source_path)
    source_digest = _sha256(source)
    if parsed.source_sha256 and parsed.source_sha256 != source_digest:
        raise PptxEditError("The template changed after inspection; inspect it again before editing")
    output = Path(output_path).expanduser().resolve(strict=False)
    if output.suffix.lower() != ".pptx":
        raise PptxEditError("The output filename must end in .pptx")
    if output == source or output.exists():
        raise PptxEditError("Choose a new output filename; existing files are never overwritten")
    root = Path(asset_root).expanduser().resolve(strict=True) if asset_root else source.parent
    prs = Presentation(str(source))
    output.parent.mkdir(parents=True, exist_ok=True)
    results = []
    with tempfile.TemporaryDirectory(prefix=".drawai-pptx-", dir=output.parent) as temp_name:
        temp = Path(temp_name)
        for index, operation in enumerate(parsed.operations):
            operation_dir = temp / str(index)
            operation_dir.mkdir()
            results.append({"op": operation.op, **_apply_operation(prs, operation, root, operation_dir)})
        draft = temp / "result.pptx"
        prs.save(str(draft))
        validation = validate_pptx_package(draft)
        # Reopening also exercises python-pptx's actual package/object reader.
        reopened = Presentation(str(draft))
        if len(reopened.slides) != len(prs.slides):
            raise PptxEditError("Exported slide count does not match the edited deck")
        if _sha256(source) != source_digest:
            raise PptxEditError("Source presentation changed while editing")
        # Atomic, exclusive publication; a concurrent output is not overwritten.
        os.link(draft, output)
    return {"schema": "drawai.pptx_edit_report.v1", "status": "ok",
            "source_pptx": str(source), "source_sha256": source_digest,
            "output_pptx": str(output), "output_sha256": _sha256(output),
            "slide_count": len(prs.slides), "operations": results,
            "package_validation": validation}


def render_pptx(source_path: str | Path, output_dir: str | Path, *,
                renderer: str | None = None, rasterizer: str | None = None) -> dict[str, Any]:
    """Render the actual PPTX through LibreOffice and Poppler, with an isolated profile."""
    source = _source(source_path)
    soffice = renderer or shutil.which("soffice")
    pdftoppm = rasterizer or shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        raise PptxEditError("PPTX preview requires LibreOffice (soffice) and Poppler (pdftoppm)")
    output = Path(output_dir).expanduser().resolve(strict=False)
    if output.exists() and any(output.iterdir()):
        raise PptxEditError("Choose an empty preview directory to avoid stale slide images")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".drawai-pptx-render-", dir=output.parent) as temp_name:
        temp = Path(temp_name)
        subprocess.run([soffice, f"-env:UserInstallation={(temp / 'profile').as_uri()}",
                        "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(source)],
                       check=True, capture_output=True, text=True, timeout=180)
        pdf = temp / f"{source.stem}.pdf"
        if not pdf.is_file():
            raise PptxEditError("LibreOffice returned without producing a PDF")
        subprocess.run([pdftoppm, "-png", "-r", "96", str(pdf), str(temp / "slide")],
                       check=True, capture_output=True, text=True, timeout=180)
        previews = sorted(temp.glob("slide-*.png"))
        if len(previews) != len(Presentation(str(source)).slides):
            raise PptxEditError("Rendered page count does not match the presentation")
        output.mkdir(exist_ok=True)
        shutil.copy2(pdf, output / pdf.name)
        for preview in previews:
            shutil.copy2(preview, output / preview.name)
    return {"schema": "drawai.pptx_preview.v1", "source_pptx": str(source),
            "pdf": str(output / pdf.name),
            "slides": [str(output / preview.name) for preview in previews]}
